"""Full-size slide downloads and a chapter/section-organized OCR Word document."""
from io import BytesIO
import json
import math
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image


class SlideExportError(Exception):
    def __init__(self, status_code, detail):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


def validate_slide_video_id(video_id):
    if not re.fullmatch(r"[A-Za-z0-9_-]+", video_id):
        raise SlideExportError(400, "Invalid video ID")


def load_download_slides(video_id, scenes_dir):
    validate_slide_video_id(video_id)
    source = Path(scenes_dir) / f"{video_id}.json"
    if not source.is_file():
        raise SlideExportError(404, "Detect slides before exporting them")
    try:
        scenes = json.loads(source.read_text(encoding="utf-8"))
    except (OSError, ValueError) as error:
        raise SlideExportError(500, f"Could not read detected slides: {error}") from error
    if not isinstance(scenes, list) or not scenes:
        raise SlideExportError(400, "No detected slides are available")
    if any(not isinstance(scene, dict) for scene in scenes):
        raise SlideExportError(409, "Invalid saved slides. Run slide detection again.")
    return scenes


def slide_seconds(scene):
    try:
        value = scene.get("time_seconds")
        if value is None:
            value = 0.0
            for component in str(scene["timestamp"]).split(":"):
                value = value * 60 + float(component)
        value = float(value)
        if not math.isfinite(value) or value < 0:
            raise ValueError
        return value
    except (KeyError, TypeError, ValueError) as error:
        raise SlideExportError(409, "Invalid slide timestamp. Run slide detection again.") from error


def slide_filename(index, scene):
    seconds = int(slide_seconds(scene))
    clock = f"{seconds // 3600:02d}-{seconds // 60 % 60:02d}-{seconds % 60:02d}"
    return f"slide_{index + 1:03d}_{clock}.jpg"


def slide_image_path(video_id, index, scenes, images_dir):
    validate_slide_video_id(video_id)
    if type(index) is not int or not 0 <= index < len(scenes):
        raise SlideExportError(404, "Slide not found")
    root = Path(images_dir).resolve()
    source = (root / video_id / f"{index}.jpg").resolve()
    if not source.is_relative_to(root) or not source.is_file():
        raise SlideExportError(409, f"Missing screenshot for slide {index + 1}. Run slide detection again.")
    return source


def build_slide_archive(video_id, indices, scenes_dir, images_dir):
    scenes = load_download_slides(video_id, scenes_dir)
    if not isinstance(indices, list) or not indices or len(indices) > len(scenes):
        raise SlideExportError(400, "Select one or more detected slides")
    if any(type(index) is not int or not 0 <= index < len(scenes) for index in indices):
        raise SlideExportError(400, "Invalid slide selection")
    indices = sorted(set(indices), key=lambda index: (slide_seconds(scenes[index]), index))
    images = [(index, slide_image_path(video_id, index, scenes, images_dir)) for index in indices]
    output = BytesIO()
    manifest = []
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for index, source in images:
            name = slide_filename(index, scenes[index])
            archive.write(source, name)
            manifest.append({
                "slide": index + 1, "time_seconds": slide_seconds(scenes[index]),
                "timestamp": scenes[index].get("timestamp"), "filename": name,
            })
        archive.writestr("screenshots.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    output.seek(0)
    return output


def clean_ocr_lines(text):
    lines = []
    for line in str(text or "").splitlines():
        line = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", line)
        line = " ".join(line.split()).lstrip("•▪●◦–—-*· ").strip()
        if line:
            lines.append(line)
    return lines


def stored_slide_lines(scene):
    blocks = []
    yolo = scene.get("yolo_detections") or {}
    if yolo.get("success", False):
        blocks.extend((item.get("bbox") or [], item.get("ocr_text"))
                      for item in yolo.get("detections", []) if item.get("ocr_text"))
    surya = scene.get("surya_ocr") or {}
    if surya.get("success", True):
        blocks.extend((item.get("bbox") or [], item.get("text"))
                      for item in surya.get("results", [])
                      if item.get("text") and not item.get("matched"))

    def order(block):
        try:
            return float(block[0][1]), float(block[0][0])
        except (TypeError, ValueError, IndexError):
            return math.inf, 0

    lines, seen = [], set()
    for _, text in sorted(blocks, key=order):
        for line in clean_ocr_lines(text):
            if line.casefold() not in seen:
                seen.add(line.casefold())
                lines.append(line)
    return lines


def recognize_slide_image(source):
    import pytesseract
    try:
        with Image.open(source) as image:
            return pytesseract.image_to_string(image, timeout=30)
    except pytesseract.TesseractNotFoundError as error:
        raise SlideExportError(
            503, "Install Tesseract or set TESSERACT_CMD, then retry the slide text export."
        ) from error
    except (OSError, RuntimeError) as error:
        raise SlideExportError(500, f"Could not extract slide text: {error}") from error


def slide_text_groups(chapters, scenes):
    """Assign every slide once, using the sidebar's half-second boundary tolerance."""
    ordered_slides = sorted(range(len(scenes)), key=lambda index: (slide_seconds(scenes[index]), index))
    chapters = sorted(chapters or [], key=slide_seconds)
    groups = []
    for chapter_number, chapter in enumerate(chapters):
        sections = chapter.get("sections") or []
        if not isinstance(sections, list) or any(not isinstance(section, dict) for section in sections):
            raise SlideExportError(409, "Invalid saved sections. Generate chapters again.")
        sections = sorted(sections, key=slide_seconds)
        if not sections:
            groups.append({"chapter": chapter, "chapter_number": chapter_number, "section": None,
                           "start": slide_seconds(chapter), "indices": []})
        for section_number, section in enumerate(sections):
            groups.append({
                "chapter": chapter, "chapter_number": chapter_number, "section": section,
                "start": min(slide_seconds(chapter), slide_seconds(section))
                         if section_number == 0 else slide_seconds(section),
                "indices": [],
            })
    if not groups:
        groups = [{"chapter": {"title": "Slides"}, "chapter_number": 0,
                   "section": None, "start": 0, "indices": []}]
    groups.sort(key=lambda group: group["start"])
    for index in ordered_slides:
        time = slide_seconds(scenes[index])
        eligible = [group for group in groups if time >= group["start"] - 0.5]
        (eligible[-1] if eligible else groups[0])["indices"].append(index)
    return groups


def new_numbered_list(document):
    numbering = document.part.numbering_part.element
    abstract_id = max([int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))] + [-1]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1.")):
        element = OxmlElement(f"w:{tag}")
        element.set(qn("w:val"), value)
        level.append(element)
    abstract.append(level)
    numbering.append(abstract)
    return numbering.add_num(abstract_id).numId


def build_slide_text_document(video_id, scenes_dir, images_dir, summaries_dir, recognize=None, transcript_language=""):
    scenes = load_download_slides(video_id, scenes_dir)
    summary = Path(summaries_dir) / f"{video_id}.json"
    if transcript_language:
        if not re.fullmatch(r"[A-Za-z0-9_-]+", transcript_language):
            raise SlideExportError(400, "Invalid transcript language")
        translated = Path(summaries_dir) / f"{video_id}_summary_{transcript_language}.json"
        if translated.is_file():
            summary = translated
    try:
        chapters = json.loads(summary.read_text(encoding="utf-8")) if summary.is_file() else []
        if not isinstance(chapters, list) or any(not isinstance(chapter, dict) for chapter in chapters):
            raise ValueError("Invalid saved chapters")
    except (OSError, ValueError) as error:
        raise SlideExportError(409, f"Could not read chapters: {error}") from error
    recognize = recognize or recognize_slide_image
    document = Document()
    document.add_heading("Slide text", 0)
    document.add_paragraph(f"Video: {video_id}. Slides are ordered by time within chapters and sections.")
    document.add_paragraph("Existing OCR text is reused; slides without text are read from their full-size screenshots.")
    previous_chapter = None
    for group in slide_text_groups(chapters, scenes):
        if previous_chapter != group["chapter_number"]:
            document.add_heading(f"{group['chapter_number'] + 1}. {group['chapter'].get('title', 'Chapter')}", 1)
            previous_chapter = group["chapter_number"]
        section = group["section"]
        if section:
            document.add_heading(str(section.get("title") or "Section"), 2)
        for index in group["indices"]:
            scene = scenes[index]
            seconds = int(slide_seconds(scene))
            clock = f"{seconds // 3600:02d}:{seconds // 60 % 60:02d}:{seconds % 60:02d}"
            document.add_heading(f"Slide {index + 1} — {clock}", 3 if section else 2)
            lines = stored_slide_lines(scene)
            if not lines:
                source = slide_image_path(video_id, index, scenes, images_dir)
                try:
                    lines = clean_ocr_lines(recognize(source))
                except SlideExportError as error:
                    raise SlideExportError(error.status_code, f"Slide {index + 1}: {error.detail}") from error
            if not lines:
                document.add_paragraph("No readable text detected on this slide.")
                continue
            number_id = new_numbered_list(document)
            for line in lines:
                paragraph = document.add_paragraph(line, style="List Number")
                properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                properties.get_or_add_ilvl().val = 0
                properties.get_or_add_numId().val = number_id
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
