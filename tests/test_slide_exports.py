import ast
import asyncio
from io import BytesIO
import json
from pathlib import Path
import shutil
from types import SimpleNamespace
import unittest
from unittest import mock
import uuid
import zipfile

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from slide_exports import (
    SlideExportError, build_slide_archive, build_slide_text_document,
    load_download_slides, slide_image_path, slide_text_groups, stored_slide_lines,
)


ROOT = Path(__file__).resolve().parents[1]


class SlideExportTests(unittest.TestCase):
    def setUp(self):
        self.root = ROOT / (".slide-export-tests-" + uuid.uuid4().hex)
        self.root.mkdir()
        self.addCleanup(shutil.rmtree, self.root)
        self.scenes_dir, self.images_dir, self.summaries_dir = [
            self.root / name for name in ("scenes", "images", "summaries")
        ]
        for directory in (self.scenes_dir, self.images_dir, self.summaries_dir):
            directory.mkdir()
        (self.images_dir / "lecture").mkdir()
        self.scenes = [
            {"timestamp": "00:00", "time_seconds": 0},
            {"timestamp": "00:10", "time_seconds": 10},
            {"timestamp": "00:20", "time_seconds": 20},
            {"timestamp": "00:30", "time_seconds": 30},
        ]
        self.chapters = [
            {"timestamp": "00:05", "title": "Opening", "sections": [
                {"timestamp": "00:05", "title": "Overview"},
                {"timestamp": "00:15", "title": "Details"},
            ]},
            {"timestamp": "00:30", "title": "Conclusion"},
        ]
        for index in range(len(self.scenes)):
            Image.new("RGB", (12, 8), (index * 30, 0, 0)).save(self.images_dir / "lecture" / f"{index}.jpg")
        self.save_scenes()
        (self.summaries_dir / "lecture.json").write_text(json.dumps(self.chapters), encoding="utf-8")

    def save_scenes(self):
        (self.scenes_dir / "lecture.json").write_text(json.dumps(self.scenes), encoding="utf-8")

    def document(self, **kwargs):
        return Document(build_slide_text_document(
            "lecture", self.scenes_dir, self.images_dir, self.summaries_dir, **kwargs,
        ))

    def test_selected_zip_has_only_requested_slides_in_time_order(self):
        output = build_slide_archive("lecture", [2, 0, 2], self.scenes_dir, self.images_dir)
        with zipfile.ZipFile(output) as archive:
            self.assertEqual(archive.namelist(), [
                "slide_001_00-00-00.jpg", "slide_003_00-00-20.jpg", "screenshots.json",
            ])
            manifest = json.loads(archive.read("screenshots.json"))
            self.assertEqual([item["slide"] for item in manifest], [1, 3])
            self.assertEqual(manifest[1]["time_seconds"], 20)
            self.assertEqual(archive.read(manifest[0]["filename"]),
                             (self.images_dir / "lecture" / "0.jpg").read_bytes())

    def test_invalid_selection_paths_and_missing_slides(self):
        for indices in (None, [], [True], [-1], [4], ["0"]):
            with self.subTest(indices=indices), self.assertRaises(SlideExportError) as result:
                build_slide_archive("lecture", indices, self.scenes_dir, self.images_dir)
            self.assertEqual(result.exception.status_code, 400)
        for video_id in ("..", "../lecture", "lecture\\other", "bad\r\nheader"):
            with self.subTest(video_id=video_id), self.assertRaises(SlideExportError):
                load_download_slides(video_id, self.scenes_dir)
        with self.assertRaises(SlideExportError) as result:
            slide_image_path("lecture", -1, self.scenes, self.images_dir)
        self.assertEqual(result.exception.status_code, 404)
        (self.images_dir / "lecture" / "2.jpg").unlink()
        with self.assertRaises(SlideExportError) as result:
            build_slide_archive("lecture", [0, 2], self.scenes_dir, self.images_dir)
        self.assertEqual(result.exception.status_code, 409)

    def test_hierarchy_assigns_each_slide_once_like_sidebar(self):
        groups = slide_text_groups(self.chapters, self.scenes)
        self.assertEqual([group["indices"] for group in groups], [[0, 1], [2], [3]])
        self.scenes[2]["time_seconds"] = 29.5
        groups = slide_text_groups(self.chapters, self.scenes)
        self.assertEqual([group["indices"] for group in groups], [[0, 1], [], [2, 3]])

    def test_ocr_document_reuses_saved_text_and_reads_remaining_slides(self):
        self.scenes[0]["yolo_detections"] = {"success": True, "detections": [
            {"bbox": [0, 20, 100, 30], "ocr_text": "• Second line\nOther detail"},
            {"bbox": [0, 0, 100, 10], "ocr_text": "Opening topic"},
        ]}
        self.scenes[0]["surya_ocr"] = {"success": True, "results": [
            {"bbox": [0, 40, 100, 50], "text": "second line"},
            {"text": "ignored matched", "matched": True},
        ]}
        self.save_scenes()
        recognize = mock.Mock(side_effect=["One\nTwo", "", "Conclusion"])
        document = self.document(recognize=recognize)
        self.assertEqual([call.args[0].name for call in recognize.call_args_list], ["1.jpg", "2.jpg", "3.jpg"])
        headings = [(p.style.name, p.text) for p in document.paragraphs if p.style.name.startswith("Heading")]
        self.assertEqual(headings, [
            ("Heading 1", "1. Opening"),
            ("Heading 2", "Overview"),
            ("Heading 3", "Slide 1 — 00:00:00"),
            ("Heading 3", "Slide 2 — 00:00:10"),
            ("Heading 2", "Details"),
            ("Heading 3", "Slide 3 — 00:00:20"),
            ("Heading 1", "2. Conclusion"),
            ("Heading 2", "Slide 4 — 00:00:30"),
        ])
        numbered = [p for p in document.paragraphs if p.style.name == "List Number"]
        self.assertEqual([p.text for p in numbered], [
            "Opening topic", "Second line", "Other detail", "One", "Two", "Conclusion",
        ])
        numbers = [p._p.pPr.numPr.numId.val for p in numbered]
        self.assertEqual(numbers[0], numbers[2])
        self.assertNotEqual(numbers[0], numbers[3], "Text numbering restarts for each slide")
        numbering = document.part.numbering_part.element
        self.assertTrue(numbering.findall(qn("w:abstractNum")))
        self.assertIn("No readable text detected on this slide.", [p.text for p in document.paragraphs])

    def test_no_chapters_and_chronological_sorting(self):
        (self.summaries_dir / "lecture.json").unlink()
        self.scenes[0]["time_seconds"], self.scenes[3]["time_seconds"] = 40, 0
        self.save_scenes()
        document = self.document(recognize=lambda _: "text")
        slides = [p.text for p in document.paragraphs if p.style.name == "Heading 2"]
        self.assertEqual(slides, [
            "Slide 4 — 00:00:00", "Slide 2 — 00:00:10", "Slide 3 — 00:00:20", "Slide 1 — 00:00:40",
        ])
        self.assertIn("1. Slides", [p.text for p in document.paragraphs])

    def test_translation_hierarchy_and_ocr_failure(self):
        (self.summaries_dir / "lecture_summary_de.json").write_text(json.dumps([
            {"timestamp": "00:00", "title": "Einleitung"},
        ]), encoding="utf-8")
        document = self.document(recognize=lambda _: "Text", transcript_language="de")
        self.assertIn("1. Einleitung", [p.text for p in document.paragraphs])
        with self.assertRaises(SlideExportError):
            self.document(recognize=lambda _: "", transcript_language="../outside")
        with self.assertRaises(SlideExportError) as result:
            self.document(recognize=mock.Mock(side_effect=SlideExportError(503, "Install Tesseract")))
        self.assertEqual(result.exception.status_code, 503)

    def test_routes_set_download_headers_and_surface_validation_errors(self):
        names = {"download_slide", "download_selected_slides", "download_slide_text_docx"}
        tree = ast.parse((ROOT / "main.py").read_text(encoding="utf-8"))
        nodes = [node for node in tree.body if isinstance(node, ast.AsyncFunctionDef) and node.name in names]
        for node in nodes:
            node.decorator_list = []
        namespace = {
            "Request": object, "asyncio": asyncio, "json": json,
            "HTTPException": SlideExportError, "SCENES_DIR": self.scenes_dir,
            "FULLSIZE_IMAGES_DIR": self.images_dir, "SUMMARIES_DIR": self.summaries_dir,
            "FileResponse": lambda source, **kwargs: SimpleNamespace(source=source, **kwargs),
            "StreamingResponse": lambda content, **kwargs: SimpleNamespace(content=b"".join(content), **kwargs),
        }
        exec(compile(ast.Module(body=nodes, type_ignores=[]), "main.py", "exec"), namespace)
        response = asyncio.run(namespace["download_slide"]("lecture", 1))
        self.assertEqual(response.filename, "slide_002_00-00-10.jpg")
        self.assertEqual(response.media_type, "image/jpeg")
        async def payload():
            return {"indices": [0, 1]}
        response = asyncio.run(namespace["download_selected_slides"]("lecture", SimpleNamespace(json=payload)))
        self.assertEqual(response.media_type, "application/zip")
        self.assertIn("attachment", response.headers["Content-Disposition"])
        with zipfile.ZipFile(BytesIO(response.content)) as archive:
            self.assertEqual(len(archive.namelist()), 3)
        with mock.patch("slide_exports.recognize_slide_image", return_value="OCR text"):
            response = asyncio.run(namespace["download_slide_text_docx"]("lecture"))
        self.assertIn(".docx", response.headers["Content-Disposition"])
        self.assertEqual(len([p for p in Document(BytesIO(response.content)).paragraphs
                              if p.text.startswith("Slide ") and p.style.name.startswith("Heading")]), 4)
        with self.assertRaises(SlideExportError) as result:
            asyncio.run(namespace["download_slide"]("lecture", 99))
        self.assertEqual(result.exception.status_code, 404)

    def test_http_downloads_and_invalid_requests(self):
        from fastapi import FastAPI, HTTPException, Request
        from fastapi.responses import FileResponse, StreamingResponse
        from fastapi.testclient import TestClient

        app = FastAPI()
        names = {"download_slide", "download_selected_slides", "download_slide_text_docx"}
        tree = ast.parse((ROOT / "main.py").read_text(encoding="utf-8"))
        nodes = [node for node in tree.body if isinstance(node, ast.AsyncFunctionDef) and node.name in names]
        namespace = {
            "app": app, "Request": Request, "asyncio": asyncio, "json": json,
            "HTTPException": HTTPException, "SCENES_DIR": self.scenes_dir,
            "FULLSIZE_IMAGES_DIR": self.images_dir, "SUMMARIES_DIR": self.summaries_dir,
            "FileResponse": FileResponse, "StreamingResponse": StreamingResponse,
        }
        exec(compile(ast.Module(body=nodes, type_ignores=[]), "main.py", "exec"), namespace)
        with TestClient(app) as client:
            response = client.get("/download_slide/lecture/0")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["content-type"], "image/jpeg")
            self.assertIn("attachment", response.headers["content-disposition"])
            self.assertEqual(response.content, (self.images_dir / "lecture" / "0.jpg").read_bytes())
            self.assertEqual(client.get("/download_slide/lecture/99").status_code, 404)
            response = client.post("/download_slides/lecture", json={"indices": [1, 0]})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["content-type"], "application/zip")
            with zipfile.ZipFile(BytesIO(response.content)) as archive:
                self.assertEqual(len(archive.namelist()), 3)
            for data in ([], {"indices": []}, {"indices": [-1]}, {"indices": [True]}):
                self.assertEqual(client.post("/download_slides/lecture", json=data).status_code, 400)
            self.assertEqual(client.post("/download_slides/lecture", content="invalid json").status_code, 400)
            with mock.patch("slide_exports.recognize_slide_image", return_value="Line 1\nLine 2"):
                response = client.get("/download_slide_text_docx/lecture")
            self.assertEqual(response.status_code, 200)
            self.assertIn("wordprocessingml.document", response.headers["content-type"])
            self.assertIn(".docx", response.headers["content-disposition"])
            self.assertEqual(len([p for p in Document(BytesIO(response.content)).paragraphs
                                  if p.style.name == "List Number"]), 8)
            with mock.patch("slide_exports.recognize_slide_image",
                            side_effect=SlideExportError(503, "Install Tesseract")):
                response = client.get("/download_slide_text_docx/lecture")
            self.assertEqual(response.status_code, 503)
            self.assertIn("Slide 1", response.json()["detail"])


if __name__ == "__main__":
    unittest.main()
