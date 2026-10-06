"""Exercise export helpers/endpoint without loading transcription or GPU models."""
import ast
import asyncio
from collections import Counter
import html
from io import BytesIO
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from types import SimpleNamespace
import unittest
from unittest import mock
import uuid
import zipfile

from docx import Document
from docx.shared import Inches
from docx.image.exceptions import UnrecognizedImageError
from docx.oxml import OxmlElement
from PIL import Image as PillowImage, UnidentifiedImageError
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image as PdfImage, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
from export_jobs import export_workspace


ROOT = Path(__file__).resolve().parents[1]


class HTTPException(Exception):
    def __init__(self, status_code, detail):
        super().__init__(detail)
        self.status_code, self.detail = status_code, detail


def load_export_functions():
    tree = ast.parse((ROOT / "main.py").read_text(encoding="utf-8"))
    names = {
        "build_outline_points", "parse_chapter_timestamp", "format_chapter_timestamp",
        "chapter_boundaries_from_topics", "chapter_boundaries_from_scenes", "export_slide_start",
        "export_sentences", "align_export_chapters", "export_key_points", "summarize_export_batch",
        "build_export_subparts", "export_archive_paths", "prepare_export_image",
        "concise_export_title", "export_thumbnail", "strip_narration", "clean_outline_title", "build_slides_pdf",
        "create_combined_chapter_documents", "collapse_word_heading", "export_chapters",
        "scene_slide_text", "chapter_slide_content", "add_word_slide_content",
    }
    nodes = [node for node in tree.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
             and node.name in names]
    for node in nodes:
        node.decorator_list = []
    namespace = dict(globals(), Request=object, FileResponse=lambda path, **kwargs: SimpleNamespace(path=path, **kwargs))
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(ROOT / "main.py"), "exec"), namespace)
    return namespace


class ExportDocumentTests(unittest.TestCase):
    def setUp(self):
        self.directory = ROOT / (".export-tests-" + uuid.uuid4().hex)
        self.directory.mkdir()
        self.addCleanup(shutil.rmtree, self.directory)
        self.functions = load_export_functions()
        for name in ("EXPORTS_DIR", "TRANSCRIPTS_DIR", "SUMMARIES_DIR", "SCENES_DIR", "VIDEO_DIR", "FULLSIZE_IMAGES_DIR"):
            directory = self.directory / name
            directory.mkdir()
            self.functions[name] = directory
        self.functions["summary_processor"] = SimpleNamespace(
            _refresh_clients=lambda: None, _provider_chain=lambda: [],
        )
        self.functions["video_processor"] = SimpleNamespace(get_video_path=lambda _: self.directory / "lecture.mp4")
        self.transcript = [
            {"start": 0, "duration": 4, "text": "Neural networks learn"},
            {"start": 4, "duration": 4, "text": "from examples. Training adjusts weights."},
            {"start": 8, "duration": 4, "text": "Validation measures generalization."},
            {"start": 12, "duration": 4, "text": "Regularization prevents overfitting."},
        ]
        self.functions["TRANSCRIPTS_DIR"].joinpath("lecture.json").write_text(json.dumps(self.transcript), encoding="utf-8")
        self.functions["SUMMARIES_DIR"].joinpath("lecture.json").write_text(json.dumps([
            {"timestamp": "00:00", "title": "Neural networks"},
            {"timestamp": "00:03", "title": "Training and evaluation"},
        ]), encoding="utf-8")
        self.functions["SCENES_DIR"].joinpath("lecture.json").write_text(json.dumps([
            {"timestamp": "00:00", "title": "Learning"},
            {"timestamp": "00:09", "title": "Validation diagram"},
        ]), encoding="utf-8")

    def run_export(self, **options):
        self.media_commands = []
        async def request_json():
            return options

        def run_media(command, **kwargs):
            self.media_commands.append(command)
            if command[0] == "ffprobe":
                return SimpleNamespace(stdout="16")
            target = Path(command[-1])
            if target.suffix == ".jpg":
                PillowImage.new("RGB", (32, 18), "white").save(target)
            else:
                target.write_bytes(b"test video")
            return SimpleNamespace(returncode=0)

        with mock.patch.object(subprocess, "run", side_effect=run_media):
            return asyncio.run(self.functions["export_chapters"]("lecture", SimpleNamespace(json=request_json)))

    def test_sentence_alignment_preserves_every_word_once(self):
        sentences = self.functions["export_sentences"](self.transcript)
        self.assertEqual(sentences[0]["text"], "Neural networks learn from examples.")
        chapters = self.functions["align_export_chapters"]([
            {"timestamp": "00:00", "title": "First"},
            {"timestamp": "00:03", "title": "Second"},
        ], sentences)
        self.assertGreater(chapters[1]["start"], 3)
        self.assertEqual(chapters[1]["start"], sentences[1]["start"])
        chunks = [[item["text"] for item in sentences if item["start"] >= part["start"]
                   and (part["end"] is None or item["start"] < part["end"])] for part in chapters]
        self.assertEqual(" ".join(text for chunk in chunks for text in chunk),
                         " ".join(item["text"] for item in self.transcript))

    def test_whisper_unpunctuated_phrases_join_until_pause_or_speaker_turn(self):
        phrases = [
            {"start": 0, "end": 2, "text": "neural networks learn"},
            {"start": 2.1, "end": 4, "text": "from examples", "speaker": "A"},
            {"start": 5.5, "duration": 2, "text": "training adjusts weights", "speaker": "A"},
            {"start": 7.6, "duration": 2, "text": "validation measures generalization", "speaker": "B"},
        ]
        sentences = self.functions["export_sentences"](phrases)
        self.assertEqual([item["text"] for item in sentences], [
            "neural networks learn from examples",
            "training adjusts weights",
            "validation measures generalization",
        ])
        parts = self.functions["align_export_chapters"]([
            {"timestamp": "00:00", "title": "Learning"},
            {"timestamp": "00:03", "title": "Evaluation"},
        ], sentences)
        self.assertEqual(parts[1]["start"], 5.5)
        self.assertEqual(" ".join(item["text"] for item in sentences),
                         " ".join(item["text"] for item in phrases))

    def test_unknown_caption_duration_does_not_invent_pauses(self):
        fragments = [
            {"start": 0, "text": "we learn"},
            {"start": 20, "text": "from examples"},
            {"start": 30, "text": "to improve predictions"},
        ]
        sentences = self.functions["export_sentences"](fragments)
        self.assertEqual(len(sentences), 1)
        self.assertEqual(sentences[0]["text"], "we learn from examples to improve predictions")

    def test_slide_content_option_adds_ocr_bullets_per_slide(self):
        def ocr(text, top):
            return {"class": "text", "bbox": [10, top, 200, top + 20], "ocr_text": text}
        self.functions["SCENES_DIR"].joinpath("lecture.json").write_text(json.dumps([
            {"timestamp": "00:00", "yolo_detections": {"success": True, "detections": [
                ocr("\u2022 Backpropagation\nGradient descent", 80), ocr("Neural Networks", 10),
                ocr("gradient descent", 120), ocr("|", 150),
            ]}},
            {"timestamp": "00:09", "surya_ocr": {"success": True, "results": [
                {"text": "Validation set", "bbox": [0, 5, 50, 20]},
                {"text": "Matched elsewhere", "bbox": [0, 30, 50, 40], "matched": True},
            ]}},
        ]), encoding="utf-8")
        response = self.run_export(include_webpage=True, include_word=True, include_pdf=True, include_outline=True, include_slide_text=True)
        with zipfile.ZipFile(response.path) as archive:
            webpage = archive.read("index.html").decode()
            word = Document(BytesIO(archive.read("chapter_document.docx")))
            outline = Document(BytesIO(archive.read("video_outline.docx")))
            self.assertGreater(len(archive.read("chapter_document.pdf")), 1000)
        self.assertIn("Validation set", [p.text for p in outline.paragraphs])
        self.assertEqual(webpage.count('<details class="slide-content"'), 2)
        self.assertIn("<ul><li>Neural Networks</li><li>Backpropagation</li><li>Gradient descent</li></ul>", webpage)
        self.assertIn("<li>Validation set</li>", webpage)
        self.assertNotIn("Matched elsewhere", webpage)
        bullets = [p.text for p in word.paragraphs if p.style.name == "List Bullet"]
        self.assertIn("Backpropagation", bullets)
        self.assertIn("Validation set", bullets)

        plain = self.run_export(include_webpage=True)
        with zipfile.ZipFile(plain.path) as archive:
            self.assertNotIn("slide-content\"", archive.read("index.html").decode())

    def test_slide_media_links_preserve_fractional_scene_time(self):
        self.functions["SCENES_DIR"].joinpath("lecture.json").write_text(json.dumps([
            {"timestamp": "00:09", "time_seconds": 9.375, "title": "Validation"},
        ]), encoding="utf-8")
        response = self.run_export(include_webpage=True, subpart_mode="slides")
        with zipfile.ZipFile(response.path) as archive:
            webpage = archive.read("index.html").decode()
            self.assertIn("slide_0000009375.jpg", archive.namelist())
        self.assertIn("#t=4.775", webpage)

    def test_clip_overlap_adds_one_second_without_changing_transcript_boundaries(self):
        response = self.run_export(include_webpage=True, include_transcripts=True)
        clips = [command for command in self.media_commands if command[-1].endswith(".mp4")]
        starts = [float(command[command.index("-ss") + 1]) for command in clips]
        self.assertEqual(starts[0], 0)
        self.assertAlmostEqual(starts[1], 4.6)
        self.assertAlmostEqual(float(clips[0][clips[0].index("-t") + 1]), 5.6)
        self.assertNotIn("-t", clips[1], "The last chapter must continue to the video end")
        with zipfile.ZipFile(response.path) as archive:
            chapters = json.loads(archive.read("chapters.json"))
            self.assertAlmostEqual(chapters[1]["start"], 5.6)
            self.assertAlmostEqual(chapters[1]["clip_start"], 4.6)
            webpage = archive.read("index.html").decode()
            self.assertNotIn("overlap", webpage)
            self.assertNotIn("Clip starts", webpage)

    def test_clip_overlap_can_be_disabled(self):
        self.run_export(include_webpage=True, clip_overlap=False)
        clips = [command for command in self.media_commands if command[-1].endswith(".mp4")]
        self.assertAlmostEqual(float(clips[1][clips[1].index("-ss") + 1]), 5.6)

    def test_clip_overlap_is_clamped_at_video_start(self):
        self.functions["TRANSCRIPTS_DIR"].joinpath("lecture.json").write_text(json.dumps([
            {"start": 0, "duration": 0.4, "text": "First sentence."},
            {"start": 0.4, "duration": 15.6, "text": "Second sentence."},
        ]), encoding="utf-8")
        self.functions["SUMMARIES_DIR"].joinpath("lecture.json").write_text(json.dumps([
            {"timestamp": "00:00", "title": "First"},
            {"timestamp": "00:00.4", "title": "Second"},
        ]), encoding="utf-8")
        self.run_export(include_webpage=True)
        clips = [command for command in self.media_commands if command[-1].endswith(".mp4")]
        self.assertEqual(len(clips), 2)
        self.assertEqual(float(clips[1][clips[1].index("-ss") + 1]), 0)

    def test_html_archive_has_real_hierarchy_and_only_required_assets(self):
        response = self.run_export(include_webpage=True, subpart_mode="both")
        with zipfile.ZipFile(response.path) as archive:
            names = archive.namelist()
            self.assertIn("index.html", names)
            self.assertTrue(all(name == "index.html" or name.endswith((".jpg", ".mp4")) for name in names), names)
            webpage = archive.read("index.html").decode()
            self.assertIn('<h2 title="Neural networks">Part 1: Neural networks</h2>', webpage)
            self.assertGreaterEqual(webpage.count("<h3 "), 3)
            self.assertIn('<details class="chapter-subparts"', webpage)
            self.assertIn('<details class="subpart"', webpage)
            self.assertNotIn("AI-generated", webpage)
            self.assertNotIn("AI not configured", webpage)
            self.assertIn('<nav class="outline"', webpage)
            self.assertRegex(webpage, r'<details class="outline-chapter"><summary><a href="#part-\d+">')
            self.assertIn('class="outline-toggle"', webpage)
            self.assertIn('href="#part-2"', webpage)
            self.assertRegex(webpage, r'<figcaption title="Neural networks">1: Neural networks</figcaption>')
            self.assertNotIn("<summary>Key point", webpage)
            self.assertIn("2: 00:00:09", webpage)
            self.assertRegex(webpage, r"<h3[^>]*>[^<]+\(00:00:\d\d\)</h3>")
            for reference in re.findall(r'(?:src|href)="([^"#][^"]*)"', webpage):
                self.assertIn(html.unescape(reference).split("#")[0], names)
            transcript_blocks = re.findall(r"<pre>(.*?)</pre>", webpage, re.DOTALL)
            self.assertEqual(" ".join(html.unescape(block).replace("\n", " ") for block in transcript_blocks),
                             " ".join(item["text"] for item in self.transcript))
        print("HTML: timestamped chapter/subpart headings, complete sentences, valid slide links; ZIP =", names)

    def test_default_export_uses_content_chapters_without_slide_sections(self):
        response = self.run_export(include_webpage=True)
        with zipfile.ZipFile(response.path) as archive:
            webpage = archive.read("index.html").decode()
            names = archive.namelist()
        self.assertIn("Part 2: Training and evaluation", webpage)
        self.assertNotIn("Slide 2:", webpage)
        self.assertFalse(any(name.startswith("slide_") for name in names), names)

    def test_explicit_transcripts_and_formats_are_preserved(self):
        response = self.run_export(include_webpage=True, include_transcripts=True,
                                   include_word=True, include_pdf=True, include_outline=True)
        with zipfile.ZipFile(response.path) as archive:
            names = archive.namelist()
            for name in ("chapter_document.docx", "chapter_document.pdf", "video_outline.docx",
                         "transcript_by_chapter.md", "chapters.json"):
                self.assertIn(name, names)
            self.assertTrue(any(name.endswith(".txt") for name in names))
            self.assertFalse(any(name.endswith(".png") for name in names))
        print("Mixed-format ZIP preserves explicitly selected documents and transcript files.")

    def test_no_format_selected_defaults_to_html(self):
        response = self.run_export()
        self.assertEqual(response.path.suffix, ".zip")
        with zipfile.ZipFile(response.path) as archive:
            self.assertIn("index.html", archive.namelist())
            self.assertFalse(any(name.endswith((".docx", ".txt", ".pdf")) for name in archive.namelist()))

    def test_selected_word_has_subpart_headings_and_full_sentences(self):
        response = self.run_export(include_word=True)
        self.assertEqual(response.path.suffix, ".docx")
        document = Document(response.path)
        headings = [paragraph.text for paragraph in document.paragraphs
                    if paragraph.style.name.startswith("Heading")]
        self.assertTrue(any("Part 1: Neural networks" in heading for heading in headings))
        self.assertTrue(any("Training adjusts weights" in heading for heading in headings))
        self.assertTrue(any("(00:00:" in heading for heading in headings))
        self.assertIn("Neural networks learn from examples.", [p.text for p in document.paragraphs])
        print("Word: verified actual document heading hierarchy =", headings)

    def test_pdf_has_actual_subpart_headings(self):
        with mock.patch.dict(self.functions, {"Paragraph": mock.Mock(wraps=Paragraph)}):
            response = self.run_export(include_pdf=True)
            paragraphs = self.functions["Paragraph"].call_args_list
        self.assertEqual(response.path.suffix, ".pdf")
        self.assertGreater(response.path.stat().st_size, 1000)
        self.assertTrue(any(call.args[1].name == "Heading2" for call in paragraphs))
        print("PDF: built", response.path.stat().st_size, "bytes with timestamped Heading2 subparts.")

    def test_slides_pdf_has_one_image_page_per_detected_slide(self):
        folder = self.functions["FULLSIZE_IMAGES_DIR"] / "lecture"
        folder.mkdir()
        for index, size in enumerate([(80, 40), (40, 80)]):
            PillowImage.new("RGB", size, "blue").save(folder / f"{index}.jpg")
        destination = self.directory / "slides.pdf"
        self.functions["build_slides_pdf"]("lecture", destination)
        pdf = destination.read_bytes()
        self.assertTrue(pdf.startswith(b"%PDF-"))
        self.assertEqual(len(re.findall(rb"/Type\s*/Page\b", pdf)), 2)
        self.assertEqual(len(re.findall(rb"/Subtype\s*/Image\b", pdf)), 2)
        self.assertRegex(pdf, rb"/MediaBox\s*\[\s*0 0 80 40\s*\]")
        self.assertRegex(pdf, rb"/MediaBox\s*\[\s*0 0 40 80\s*\]")
        (folder / "1.jpg").unlink()
        with self.assertRaises(HTTPException) as raised:
            self.functions["build_slides_pdf"]("lecture", destination)
        self.assertEqual(raised.exception.status_code, 409)

    def test_compact_export_thumbnails_and_full_title_on_hover(self):
        title = "A lengthy title describing neural networks and many learning examples"
        self.functions["SUMMARIES_DIR"].joinpath("lecture.json").write_text(json.dumps([
            {"timestamp": "00:00", "title": title},
        ]), encoding="utf-8")
        response = self.run_export(include_webpage=True)
        with zipfile.ZipFile(response.path) as archive:
            webpage = archive.read("index.html").decode()
        self.assertIn(f'title="{title}"', webpage)
        self.assertLessEqual(len(self.functions["concise_export_title"](title)), 36)
        self.assertIn('class="chapter-heading"', webpage)
        self.assertIn('class="subpart-heading"', webpage)
        self.assertIn('.export-thumbnail img{width:120px;height:68px', webpage)
        self.assertIn('target="_blank" rel="noopener" title="Enlarge:', webpage)

    def test_slides_pdf_http_download_and_temporary_cleanup(self):
        from fastapi import FastAPI, HTTPException as FastAPIException
        from fastapi.responses import FileResponse
        from fastapi.testclient import TestClient
        from starlette.background import BackgroundTask
        from export_jobs import cleanup_export_workspace
        folder = self.functions["FULLSIZE_IMAGES_DIR"] / "lecture"
        folder.mkdir()
        for index in range(2):
            PillowImage.new("RGB", (40, 20), "blue").save(folder / f"{index}.jpg")
        app = FastAPI()
        namespace = dict(self.functions, app=app, FileResponse=FileResponse,
                         HTTPException=FastAPIException, BackgroundTask=BackgroundTask,
                         cleanup_export_workspace=cleanup_export_workspace)
        tree = ast.parse((ROOT / "main.py").read_text(encoding="utf-8"))
        route = next(node for node in tree.body if isinstance(node, ast.AsyncFunctionDef)
                     and node.name == "download_slides_pdf")
        builder = next(node for node in tree.body if isinstance(node, ast.FunctionDef)
                       and node.name == "build_slides_pdf")
        exec(compile(ast.Module(body=[builder, route], type_ignores=[]), str(ROOT / "main.py"), "exec"), namespace)
        with TestClient(app) as client:
            response = client.get("/download_slides_pdf/lecture")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["content-type"], "application/pdf")
            self.assertTrue(response.content.startswith(b"%PDF-"))
            self.assertEqual(client.get("/download_slides_pdf/missing").status_code, 404)

    def test_saved_slides_are_copied_not_reextracted(self):
        folder = self.functions["FULLSIZE_IMAGES_DIR"] / "lecture"
        folder.mkdir()
        for index in range(2):
            PillowImage.new("RGB", (40, 20), "blue").save(folder / f"{index}.jpg")
        response = self.run_export(include_webpage=True, subpart_mode="both")
        image_times = [float(command[command.index("-ss") + 1])
                       for command in self.media_commands if command[-1].endswith(".jpg")]
        self.assertNotIn(0, image_times)
        self.assertNotIn(9, image_times)
        self.assertEqual(image_times, [], "Chapter images reuse the slide on screen at the chapter start")
        with zipfile.ZipFile(response.path) as archive:
            self.assertEqual(archive.read("slide_0000000000.jpg"), (folder / "0.jpg").read_bytes())
            self.assertEqual(archive.read("slide_0000009000.jpg"), (folder / "1.jpg").read_bytes())
            self.assertEqual(archive.read("01_Neural_networks.jpg"), (folder / "0.jpg").read_bytes())
            self.assertEqual(archive.read("02_Training_and_evaluation.jpg"), (folder / "0.jpg").read_bytes())

    def test_missing_saved_images_are_extracted_only_once_per_timestamp(self):
        self.run_export(include_webpage=True, subpart_mode="both")
        image_times = [float(command[command.index("-ss") + 1])
                       for command in self.media_commands if command[-1].endswith(".jpg")]
        self.assertEqual(image_times.count(0), 1, "Chapter/slide images at the same time should share a capture")
        self.assertEqual(image_times.count(9), 1)

    def test_packaged_export_survives_locked_intermediate_files(self):
        paths = set()
        def locked(path):
            paths.add(Path(path))
            raise OSError(145, "The directory is not empty")
        try:
            with mock.patch("export_jobs.shutil.rmtree", side_effect=locked), \
                    mock.patch("export_jobs.time.sleep"), self.assertLogs("export_jobs", level="WARNING"):
                response = self.run_export(include_webpage=True)
            with zipfile.ZipFile(response.path) as archive:
                self.assertIn("index.html", archive.namelist())
                self.assertIsNone(archive.testzip())
        finally:
            for path in paths:
                shutil.rmtree(path)

    def test_ai_points_use_specific_titles_and_errors_are_not_hidden(self):
        processor = self.functions["summary_processor"]
        processor._provider_chain = lambda: [("gemini", "test-model")]
        calls = []

        def complete(*args):
            calls.append(args)
            return json.dumps({
                "summary": "Neural networks learn patterns from examples.",
                "points": [{
                    "sentence_id": 0, "title": "Learning from examples",
                    "text": "Neural networks learn patterns from examples.",
                }],
            })

        processor._complete = complete
        sentences = self.functions["export_sentences"](self.transcript)
        points, method, summary = asyncio.run(self.functions["export_key_points"](sentences))
        self.assertEqual(points[0]["title"], "Learning from examples")
        self.assertEqual(method, "AI-summarized key points")
        self.assertEqual(summary, "Neural networks learn patterns from examples.")
        self.assertEqual(len(calls), 1)
        processor._complete = lambda *args: json.dumps({
            "summary": "A valid short summary.",
            "points": [{"sentence_id": 0, "title": "Key point", "text": "Fake."}],
        })
        with self.assertRaises(HTTPException) as raised:
            asyncio.run(self.functions["export_key_points"](sentences))
        self.assertEqual(raised.exception.status_code, 502)
        self.assertIn("generic", raised.exception.detail)

    def test_narration_fillers_are_removed_and_ai_titles_are_used(self):
        strip = self.functions["strip_narration"]
        self.assertEqual(strip("The lecture explains that gradient descent minimizes the loss."),
                         "Gradient descent minimizes the loss.")
        self.assertEqual(strip("The speaker points out that grades are weighted."), "Grades are weighted.")
        self.assertEqual(strip("In this section, the instructor discusses regularization methods for deep networks."),
                         "Regularization methods for deep networks.")
        self.assertEqual(strip("Validation measures generalization."), "Validation measures generalization.")
        self.assertEqual(self.functions["clean_outline_title"]("Navigating to assignments and ..."),
                         "Navigating to assignments and")
        processor = self.functions["summary_processor"]
        processor._provider_chain = lambda: [("gemini", "test-model")]
        prompts = []

        def complete(provider, model, prompt, temperature):
            prompts.append(prompt)
            return json.dumps({
                "title": "Learning From Examples",
                "summary": "The lecture explains that networks learn from examples.",
                "points": [{"sentence_id": 0, "title": "Learning", "text": "The speaker notes that data matters."}],
                "section_titles": {"2": "Validation Diagram"},
            })

        processor._complete = complete
        response = self.run_export(include_webpage=True, subpart_mode="both")
        with zipfile.ZipFile(response.path) as archive:
            webpage = archive.read("index.html").decode()
        self.assertIn("Never mention the lecture, speaker", prompts[0])
        self.assertIn("Part 1: Learning From Examples</h2>", webpage)
        self.assertIn("Networks learn from examples.", webpage)
        self.assertIn("Data matters.", webpage)
        self.assertNotIn("The lecture explains", webpage)
        self.assertIn("Validation Diagram", webpage)

    def test_long_chapters_keep_bounded_batches_and_short_combined_summary(self):
        processor = self.functions["summary_processor"]
        processor._provider_chain = lambda: [("groq", "test-model")]
        calls = []

        def complete(*args):
            calls.append(args)
            return json.dumps({
                "summary": "The chapter explains learning.",
                "points": [{"sentence_id": 0, "title": "Learning", "text": "Learning uses examples."}],
            })

        processor._complete = complete
        sentences = [{"start": index, "text": "Examples " * 700 + "."} for index in range(3)]
        points, _, summary = asyncio.run(self.functions["export_key_points"](sentences))
        self.assertEqual(len(calls), 4)
        self.assertEqual([point["start"] for point in points], [0, 1, 2])
        self.assertEqual(summary, "The chapter explains learning.")
        self.assertTrue(all(len(call[2]) < 11000 for call in calls))

    def test_one_ai_request_per_chapter_supplies_html_word_and_pdf_summaries(self):
        processor = self.functions["summary_processor"]
        processor._provider_chain = lambda: [("gemini", "test-model")]
        calls = []

        def complete(*args):
            calls.append(args)
            return json.dumps({
                "summary": "This chapter teaches neural network fundamentals.",
                "points": [{
                    "sentence_id": 0,
                    "title": "Neural network fundamentals",
                    "text": "Neural networks learn patterns from examples.",
                }],
            })

        processor._complete = complete
        paragraph = mock.Mock(wraps=Paragraph)
        with mock.patch.dict(self.functions, {"Paragraph": paragraph}):
            response = self.run_export(
                include_webpage=True, include_word=True, include_pdf=True
            )

        with zipfile.ZipFile(response.path) as archive:
            webpage = archive.read("index.html").decode()
            document = Document(BytesIO(archive.read("chapter_document.docx")))
        self.assertEqual(len(calls), 2)
        self.assertEqual(webpage.count("This chapter teaches neural network fundamentals."), 2)
        self.assertEqual(webpage.count('<details class="chapter-subparts"'), 2)
        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertEqual(
            document_text.count(
                "Summary: This chapter teaches neural network fundamentals."
            ),
            2,
        )
        self.assertTrue(any(
            "<b>Summary:</b> This chapter teaches neural network fundamentals."
            in str(call.args[0])
            for call in paragraph.call_args_list
        ))

    def test_scorm_contains_manifest_without_internal_transcript_files(self):
        response = self.run_export(include_scorm=True)
        with zipfile.ZipFile(response.path) as archive:
            names = archive.namelist()
            self.assertIn("imsmanifest.xml", names)
            self.assertIn("scorm_api.js", names)
            self.assertFalse(any(name.endswith((".txt", ".png", ".json")) for name in names))
            manifest = archive.read("imsmanifest.xml").decode()
            for reference in re.findall(r'<file href="([^"]+)"', manifest):
                self.assertIn(reference, names)

    def test_images_and_clips_work_without_any_transcript(self):
        self.functions["TRANSCRIPTS_DIR"].joinpath("lecture.json").unlink()
        response = self.run_export(include_images=True, include_clips=True)
        with zipfile.ZipFile(response.path) as archive:
            names = archive.namelist()
            self.assertTrue(any(name.endswith(".jpg") for name in names))
            self.assertTrue(any(name.endswith(".mp4") for name in names))
            self.assertTrue(all(name.endswith((".jpg", ".mp4")) for name in names))

    def test_fixed_intervals_and_part_timestamp_option_remain_selectable(self):
        response = self.run_export(include_webpage=True, interval_minutes=0.1,
                                   subpart_mode="slides", timestamp_mode="part")
        with zipfile.ZipFile(response.path) as archive:
            webpage = archive.read("index.html").decode()
        self.assertIn("Interval 2", webpage)
        self.assertIn("2: 00:00:09", webpage)
        self.assertNotRegex(webpage, r"<h3>[^<]+\(00:00:")


if __name__ == "__main__":
    unittest.main()
