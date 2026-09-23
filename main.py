from fastapi import FastAPI, Request, Header, HTTPException, BackgroundTasks, UploadFile, File, APIRouter
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
import os
import json
import hashlib
from typing import Optional, List, Dict, Any, Union
from pathlib import Path
import asyncio
from concurrent.futures import ThreadPoolExecutor
import glob
import subprocess
import tempfile
import zipfile
import shutil
import traceback
import yt_dlp
from docx import Document
from docx.shared import Inches
from docx.image.exceptions import UnrecognizedImageError
from docx.oxml import OxmlElement
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image as PdfImage, Paragraph, SimpleDocTemplate, Spacer
from youtube_transcript_api import YouTubeTranscriptApi
import re
import html
from collections import Counter
from PIL import Image as PillowImage, UnidentifiedImageError
from dotenv import load_dotenv

from project_paths import STATIC_DIR, VIDEO_DIR, TRANSCRIPTS_DIR, SCENES_DIR, THUMBNAILS_DIR, FULLSIZE_IMAGES_DIR, SUMMARIES_DIR, EXPORTS_DIR, DETECTIONS_DIR, OCR_RESULTS_DIR, ensure_app_directories

# Import processing modules
from processors.video_processor import VideoProcessor
from processors.scene_processor import SceneProcessor
from processors.ocr_processor import OCRProcessor
from processors.transcript_processor import TranscriptProcessor
from processors.embedding_processor import EmbeddingProcessor
from processors.summary_processor import SummaryProcessor

load_dotenv()

# Initialize FastAPI app
app = FastAPI()

ensure_app_directories()

@app.middleware("http")
async def disable_frontend_cache(request: Request, call_next):
    response = await call_next(request)
    if request.url.path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
    return response

@app.get("/yolo_status")
async def yolo_status():
    """Return the current availability of the YOLO model."""
    try:
        OCRProcessor.initialize_models()

        loaded = OCRProcessor.yolo_model is not None

        return JSONResponse({
            "loaded": loaded,
            "error": None if loaded else "YOLOv8 model not loaded"
        })

    except Exception as e:
        return JSONResponse({
            "loaded": False,
            "error": str(e)
        })

@app.get("/runtime_status")
async def runtime_status():
    """Report the compute device used by Whisper and available CUDA GPU."""
    summary_processor._refresh_gemini_client()
    summary_processor._refresh_openai_client()
    translation_status = {
        "translation_default_provider": (
            "Gemini" if summary_processor.model
            else "OpenAI" if summary_processor.openai_client
            else None
        ),
        "translation_default_model": (
            summary_processor.model_name
            if summary_processor.model
            else summary_processor.openai_model_name if summary_processor.openai_client
            else None
        ),
        "gemini_configured": bool(summary_processor.model),
        "openai_configured": bool(summary_processor.openai_client),
        "huggingface_token_configured": bool(
            os.getenv("HUGGINGFACE_TOKEN") or os.getenv("HF_TOKEN")
        ),
    }
    try:
        import torch

        cuda_available = bool(torch.cuda.is_available())
        gpu_name = torch.cuda.get_device_name(0) if cuda_available else None
        return JSONResponse({
            "cuda_available": cuda_available,
            "gpu_name": gpu_name,
            "cuda_version": torch.version.cuda,
            "whisper_device": "cuda" if cuda_available else "cpu",
            "whisper_compute_type": "float16" if cuda_available else "int8",
            "colab_gpu": os.getenv("COLAB_GPU") or None,
            **translation_status,
        })
    except Exception as error:
        return JSONResponse({
            "cuda_available": False,
            "gpu_name": None,
            "cuda_version": None,
            "whisper_device": "cpu",
            "whisper_compute_type": "int8",
            "colab_gpu": os.getenv("COLAB_GPU") or None,
            "error": str(error),
            **translation_status,
        })

# Mount static directory
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
templates = Jinja2Templates(directory=str(Path(__file__).resolve().parent / "templates"))

# Initialize processors
video_processor = VideoProcessor()
scene_processor = SceneProcessor()
transcript_processor = TranscriptProcessor()
embedding_processor = EmbeddingProcessor()
summary_processor = SummaryProcessor()

def parse_uploaded_transcript(text, suffix):
    if suffix == ".json":
        data = json.loads(text)
        return data if isinstance(data, list) else data.get("segments", [])
    lines = text.replace("\r\n", "\n").split("\n")
    result = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if "-->" in line:
            start = line.split("-->")[0].strip().replace(",", ".")
            parts = start.split(":")
            seconds = sum(float(value) * (60 ** position) for position, value in enumerate(reversed(parts)))
            index += 1
            text_lines = []
            while index < len(lines) and lines[index].strip():
                text_lines.append(lines[index].strip())
                index += 1
            if text_lines:
                result.append({"start": seconds, "duration": 0, "text": " ".join(text_lines)})
        elif line and not line.isdigit():
            result.append({"start": result[-1]["start"] + 1 if result else 0, "duration": 0, "text": line})
        index += 1
    return result

def detect_transcript_language(transcript):
    sample = " ".join(item.get("text", "") for item in transcript[:30])
    if re.search(r"[\u0600-\u06ff]", sample):
        return "ar"
    words = set(re.findall(r"[A-Za-zÀ-ÿ]+", sample.lower()))
    if words & {"der", "die", "das", "und", "ist"}:
        return "de"
    if words & {"jest", "nie", "oraz", "dla"}:
        return "pl"
    return "en"

def build_outline_points(transcript, max_points=4):
    """Select concise, timestamped key points for the Word outline."""
    candidates = []
    for item in transcript:
        text = " ".join(str(item.get("text", "")).split())
        if not text:
            continue
        for sentence in re.split(r"(?<=[.!?])\s+", text):
            sentence = sentence.strip()
            words = re.findall(r"[A-Za-zÀ-ÿ\u0600-\u06ff]{3,}", sentence.lower())
            if sentence and words:
                title_words = sentence.split()[:8]
                candidates.append({
                    "start": float(item.get("start", 0)),
                    "text": sentence,
                    "title": " ".join(title_words).rstrip(".,;:!?") or "Key point",
                    "words": words,
                })
    if not candidates:
        return []
    frequencies = Counter(
        word for candidate in candidates for word in candidate["words"]
    )
    for candidate in candidates:
        candidate["score"] = sum(
            frequencies[word] for word in set(candidate["words"])
        ) / max(1, len(candidate["words"]))
    selected = sorted(
        candidates,
        key=lambda candidate: (-candidate["score"], candidate["start"]),
    )[:max_points]
    return sorted(selected, key=lambda candidate: candidate["start"])

@app.post("/upload_transcript/{video_id}")
async def upload_transcript(video_id: str, transcript: UploadFile = File(...)):
    text = (await transcript.read()).decode("utf-8-sig")
    transcript_data = parse_uploaded_transcript(text, Path(transcript.filename or "").suffix.lower())
    if not transcript_data:
        raise HTTPException(status_code=400, detail="No transcript segments found")
    output_path = TRANSCRIPTS_DIR / f"{video_id}_uploaded.json"
    output_path.write_text(json.dumps(transcript_data, ensure_ascii=False), encoding="utf-8")
    return {"success": True, "transcript": transcript_data, "source_language": detect_transcript_language(transcript_data)}

@app.post("/translate_transcript/{video_id}")
async def translate_transcript(video_id: str, request: Request):
    data = await request.json()
    transcript = data.get("transcript", [])
    target = data.get("target_language")
    requested_model = data.get("model") or None
    if target not in {"de", "en", "ar", "pl"}:
        raise HTTPException(status_code=400, detail="Unsupported translation language")
    if requested_model not in {None, "gemini-3.6-flash", "gemini-2.5-flash", "gpt-4.1-mini"}:
        raise HTTPException(status_code=400, detail="Unsupported translation model")
    if (
        not isinstance(transcript, list)
        or not transcript
        or any(not isinstance(item, dict) for item in transcript)
    ):
        raise HTTPException(status_code=400, detail="No transcript segments were provided")

    chapters = []
    summary_path = SUMMARIES_DIR / f"{video_id}.json"
    if summary_path.is_file():
        try:
            chapters = json.loads(summary_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as error:
            raise HTTPException(
                status_code=500,
                detail=f"Could not read chapter summary: {error}",
            )
        if not isinstance(chapters, list):
            chapters = []
        if chapters and isinstance(chapters[0], list):
            chapters = [
                chapter
                for chapter_group in chapters
                if isinstance(chapter_group, list)
                for chapter in chapter_group
            ]
    chapter_segments = [
        {
            "start": index,
            "duration": 0,
            "text": str(chapter.get("title", "")),
        }
        for index, chapter in enumerate(chapters)
        if isinstance(chapter, dict) and str(chapter.get("title", "")).strip()
    ]
    try:
        translation = await summary_processor.translate_transcript(
            [*transcript, *chapter_segments], target, requested_model
        )
    except (RuntimeError, ValueError, asyncio.TimeoutError, json.JSONDecodeError) as error:
        raise HTTPException(status_code=502, detail=str(error)) from error
    translated = translation["transcript"][:len(transcript)]
    translated_title_segments = translation["transcript"][len(transcript):]
    (TRANSCRIPTS_DIR / f"{video_id}_translated_{target}.json").write_text(
        json.dumps(translated, ensure_ascii=False), encoding="utf-8"
    )
    translated_chapters = []
    title_by_index = {
        int(item.get("start", index)): str(item.get("text", "")).strip()
        for index, item in enumerate(translated_title_segments)
        if isinstance(item, dict) and str(item.get("text", "")).strip()
    }
    for index, chapter in enumerate(chapters):
        if not isinstance(chapter, dict):
            continue
        translated_chapter = dict(chapter)
        if index in title_by_index:
            translated_chapter["title"] = title_by_index[index]
        translated_chapters.append(translated_chapter)
    if translated_chapters:
        (SUMMARIES_DIR / f"{video_id}_summary_{target}.json").write_text(
            json.dumps(translated_chapters, ensure_ascii=False), encoding="utf-8"
        )
    return {
        "success": True,
        "transcript": translated,
        "chapters": translated_chapters,
        "language": target,
        "provider": translation["provider"],
        "model": translation["model"],
        "batch_count": translation["batch_count"],
        "elapsed_seconds": translation["elapsed_seconds"],
    }

# Create thread pool for background tasks
executor = ThreadPoolExecutor(max_workers=2)

# Dictionary to store SSE clients by video_id
sse_clients = {}

async def send_sse_update(video_id, event_data):
    """Send an SSE update to all clients for a specific video."""
    if video_id in sse_clients:
        data_str = json.dumps(event_data)
        for queue in sse_clients[video_id]:
            try:
                await queue.put(data_str)
            except Exception as e:
                print(f"Error sending SSE update: {str(e)}")

# Initialize OCR processor with SSE update function
ocr_processor = OCRProcessor(send_sse_update=send_sse_update)

# Routes

@app.get("/", response_class=HTMLResponse)
async def root(request: Request):
    response = templates.TemplateResponse(
    request=request,
    name="index.html"
    )
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    return response

@app.post("/upload_video")
async def upload_video(background_tasks: BackgroundTasks, video: UploadFile = File(...)):
    temporary_path = None
    try:
        # Hash and persist the upload in chunks to avoid holding large videos in RAM.
        hasher = hashlib.sha256()
        with tempfile.NamedTemporaryFile(
            mode="wb",
            dir=str(VIDEO_DIR),
            prefix=".upload-",
            suffix=".tmp",
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
            while True:
                chunk = await video.read(1024 * 1024)
                if not chunk:
                    break
                hasher.update(chunk)
                temporary_file.write(chunk)

        video_hash = hasher.hexdigest()
        video_path = str(VIDEO_DIR / f"{video_hash}.mp4")
        metadata_path = VIDEO_DIR / "metadata.json"
        metadata = {}
        if metadata_path.exists():
            try:
                metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                metadata = {}
        metadata[video_hash] = {"filename": video.filename or f"{video_hash}.mp4"}
        metadata_path.write_text(json.dumps(metadata), encoding="utf-8")

        # Check if video exists
        if os.path.exists(video_path):
            print("Video exists: ", os.path.exists(video_path))
            return await video_processor.handle_existing_video(video_hash, video_path, background_tasks)

        # Move the completed upload into its content-addressed location.
        os.replace(str(temporary_path), video_path)
        temporary_path = None

        return await video_processor.process_new_video(video_hash, video_path, background_tasks)

    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)})
    finally:
        if temporary_path and temporary_path.exists():
            temporary_path.unlink()

@app.get("/uploaded_videos")
async def uploaded_videos():
    metadata_path = VIDEO_DIR / "metadata.json"
    try:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8")) if metadata_path.exists() else {}
    except (OSError, json.JSONDecodeError):
        metadata = {}
    videos = []
    for video_path in sorted(VIDEO_DIR.glob("*"), key=lambda path: path.stat().st_mtime, reverse=True):
        if (
            not video_path.is_file()
            or video_path.name == "metadata.json"
            or video_path.name.endswith("_video_quality.json")
        ):
            continue
        videos.append({
            "video_id": video_path.stem,
            "filename": metadata.get(video_path.stem, {}).get("filename", video_path.name),
            "size_bytes": video_path.stat().st_size,
            "modified_at": video_path.stat().st_mtime
        })
    return {"success": True, "videos": videos}

@app.get("/uploaded_videos/{video_id}")
async def load_uploaded_video(video_id: str, background_tasks: BackgroundTasks):
    video_path = video_processor.get_video_path(video_id)
    if not video_path:
        raise HTTPException(status_code=404, detail="Uploaded video not found")
    return await video_processor.handle_existing_video(video_id, video_path, background_tasks)

@app.get("/video/{video_id}")
async def stream_video(video_id: str, range: Optional[str] = Header(None)):
    try:
        video_path = video_processor.get_video_path(video_id)
        if not video_path:
            raise HTTPException(status_code=404, detail="Video not found")
        return await video_processor.stream_video(video_path, range)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/scenes/{video_id}")
async def get_scenes(video_id: str):
    return await scene_processor.get_scenes(video_id)

@app.post("/detect_scenes/{video_id}")
async def detect_scenes(video_id: str, request: Request, background_tasks: BackgroundTasks):
    video_path = video_processor.get_video_path(video_id)
    if not video_path:
        raise HTTPException(status_code=404, detail="Video not found")
    data = await request.json()
    mode = data.get("mode", "adaptive")
    try:
        options = {
            "adaptive_detail": str(data.get("adaptive_detail", "balanced")),
            "content_threshold": float(data.get("content_threshold", 27)),
            "minimum_slide_duration": float(data.get("minimum_slide_duration", 10)),
            "maximum_slides_per_hour": int(data.get("maximum_slides_per_hour", 60)),
            "merge_similar_slides": bool(data.get("merge_similar_slides", True)),
            "include_chapter_boundaries": bool(data.get("include_chapter_boundaries", True)),
            "screenshot_height": int(data.get("screenshot_height", 720)),
        }
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Invalid scene detection options")
    if mode not in {"adaptive", "content", "chapters"}:
        raise HTTPException(status_code=400, detail="Invalid scene detection mode")
    if options["adaptive_detail"] not in {"fewer", "balanced", "more"}:
        raise HTTPException(status_code=400, detail="Invalid adaptive detail level")
    if not 10 <= options["content_threshold"] <= 60:
        raise HTTPException(status_code=400, detail="Content-cut threshold must be between 10 and 60")
    if options["minimum_slide_duration"] not in {5, 10, 20, 30}:
        raise HTTPException(status_code=400, detail="Invalid minimum slide duration")
    if options["maximum_slides_per_hour"] not in {0, 40, 60, 90}:
        raise HTTPException(status_code=400, detail="Invalid maximum slides per hour")
    if options["screenshot_height"] not in {480, 720, 1080}:
        raise HTTPException(status_code=400, detail="Invalid screenshot quality")
    print(f"Starting scene detection for {video_id}: mode={mode}, options={options}")
    await scene_processor.start_scene_detection(video_id, video_path, background_tasks, mode, options)
    return JSONResponse({"success": True, "message": "Scene detection started"})

@app.get("/scene_detections/{video_id}/{scene_index}")
async def get_scene_detections(video_id: str, scene_index: int):
    return await scene_processor.get_scene_detections(video_id, scene_index)

@app.get("/thumbnails/{video_id}/{filename}")
async def get_thumbnail(video_id: str, filename: str):
    base_dir = THUMBNAILS_DIR.resolve()
    thumbnail_path = (base_dir / video_id / filename).resolve()
    if not thumbnail_path.is_relative_to(base_dir) or not thumbnail_path.is_file():
        raise HTTPException(status_code=404, detail="Thumbnail not found")
    return FileResponse(thumbnail_path)

@app.get("/fullsize_images/{video_id}/{filename}")
async def get_fullsize_image(video_id: str, filename: str):
    base_dir = FULLSIZE_IMAGES_DIR.resolve()
    image_path = (base_dir / video_id / filename).resolve()
    if not image_path.is_relative_to(base_dir) or not image_path.is_file():
        raise HTTPException(status_code=404, detail="Image not found")
    return FileResponse(image_path)

@app.get("/download_scene_screenshots/{video_id}")
async def download_scene_screenshots(video_id: str):
    scene_path = SCENES_DIR / f"{video_id}.json"
    if not scene_path.is_file():
        raise HTTPException(status_code=404, detail="Run slide detection before downloading screenshots")

    try:
        scenes = json.loads(scene_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise HTTPException(status_code=500, detail=f"Could not read detected slides: {error}")
    if not isinstance(scenes, list) or not scenes:
        raise HTTPException(status_code=400, detail="No detected slide screenshots are available")

    image_dir = FULLSIZE_IMAGES_DIR / video_id
    with tempfile.TemporaryDirectory(dir=EXPORTS_DIR) as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        manifest = []
        image_count = 0
        for index, scene in enumerate(scenes):
            source = image_dir / f"{index}.jpg"
            if not source.is_file():
                continue
            filename = f"slide_{index + 1:03d}_{scene.get('timestamp', 'unknown').replace(':', '-')}.jpg"
            (temp_dir / filename).write_bytes(source.read_bytes())
            manifest.append({
                "slide": index + 1,
                "timestamp": scene.get("timestamp"),
                "time_seconds": scene.get("time_seconds"),
                "filename": filename
            })
            image_count += 1

        if not image_count:
            raise HTTPException(status_code=400, detail="Detected slides have no generated screenshots")

        (temp_dir / "screenshots.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        zip_path = EXPORTS_DIR / f"{video_id}_slide_screenshots.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for file_path in temp_dir.iterdir():
                archive.write(file_path, file_path.name)

    return FileResponse(
        zip_path,
        media_type="application/zip",
        filename=f"{video_id}_slide_screenshots.zip"
    )

@app.get("/ocr_text/{video_id}")
async def get_ocr_text(video_id: str):
    return await ocr_processor.get_ocr_text(video_id)

@app.post("/stop_ocr/{video_id}")
async def stop_ocr(video_id: str):
    return await ocr_processor.stop_ocr(video_id)

@app.post("/start_ocr/{video_id}")
async def start_ocr(video_id: str):
    return await ocr_processor.start_ocr(video_id)

@app.post("/process_surya_ocr/{video_id}")
async def process_surya_ocr(video_id: str):
    return await ocr_processor.process_surya_ocr(video_id)

@app.post("/set_ocr_preference")
async def set_ocr_preference(request: Request):
    data = await request.json()
    return await ocr_processor.set_preference(data.get("preference", "tesseract"))

@app.get("/get_ocr_preference")
def get_ocr_preference():
    return {"preference": ocr_processor.get_preference()}

@app.get("/get_transcript/{video_id}/{source}")
async def get_transcript(video_id: str, source: str):
    return await transcript_processor.get_transcript(video_id, source)

@app.post("/generate_whisper_transcript/{video_id}")
async def generate_whisper_transcript(video_id: str, request: Request, background_tasks: BackgroundTasks):
    try:
        options = await request.json()
    except json.JSONDecodeError:
        options = {}
    return await transcript_processor.generate_whisper_transcript(
        video_id, background_tasks, options.get("model", "turbo"), options.get("prompt"),
        bool(options.get("diarization", False)), bool(options.get("force", False))
    )

@app.get("/whisper_transcript_status/{video_id}")
async def whisper_transcript_status(video_id: str):
    return await transcript_processor.get_whisper_status(video_id)

@app.post("/speaker_names/{video_id}")
async def speaker_names(video_id: str, request: Request):
    data = await request.json()
    names = data.get("names")
    if not isinstance(names, dict):
        raise HTTPException(status_code=400, detail="Speaker names must be an object")
    return await transcript_processor.save_speaker_names(video_id, names)

@app.post("/compute_embeddings/{video_id}")
async def compute_embeddings_endpoint(video_id: str, background_tasks: BackgroundTasks):
    return await embedding_processor.compute_embeddings(video_id, background_tasks)

@app.get("/embeddings_status/{video_id}")
async def embeddings_status(video_id: str):
    return await embedding_processor.get_status(video_id)

@app.get("/get_transcript_ocr_relationships/{video_id}")
async def get_transcript_ocr_relationships(video_id: str):
    return await embedding_processor.get_relationships(video_id)

@app.get("/find_ocr_for_transcript/{video_id}/{transcript_index}")
async def get_ocr_for_transcript(video_id: str, transcript_index: int):
    return await embedding_processor.find_ocr_for_transcript(video_id, transcript_index)

@app.get("/find_transcript_for_ocr/{video_id}/{scene_index}")
async def get_transcript_for_ocr(video_id: str, scene_index: int, ocr_text: str):
    return await embedding_processor.find_transcript_for_ocr(video_id, scene_index, ocr_text)

@app.get("/find_scene_for_transcript/{video_id}/{transcript_index}")
async def get_scene_for_transcript(video_id: str, transcript_index: int):
    return await embedding_processor.find_scene_for_transcript(video_id, transcript_index)

# SSE routes and handlers
@app.get("/ocr_progress/{video_id}")
async def ocr_progress(video_id: str):
    async def event_generator():
        if video_id not in sse_clients:
            sse_clients[video_id] = []
        
        queue = asyncio.Queue()
        sse_clients[video_id].append(queue)
        
        try:
            await queue.put(json.dumps({
                "event": "connected",
                "data": {"message": "SSE connection established"}
            }))
            
            while True:
                data = await queue.get()
                yield f"data: {data}\n\n"
        except asyncio.CancelledError:
            pass
        finally:
            if video_id in sse_clients and queue in sse_clients[video_id]:
                sse_clients[video_id].remove(queue)
                if not sse_clients[video_id]:
                    del sse_clients[video_id]
    
    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )

def extract_video_id(url):
    """Extract video ID from YouTube URL."""
    pattern = r'(?:v=|\/)([0-9A-Za-z_-]{11}).*'
    match = re.search(pattern, url)
    return match.group(1) if match else None

def _move_moov_atom_to_front(video_path: Path) -> None:
    """Remux an mp4 in place so its moov atom is at the front (faststart).

    yt-dlp's ffmpeg merge does not add ``+faststart`` by default. Without it, a
    merged mp4 can store its moov atom at the end of the file, so a browser's
    <video> element must fetch close to the end of the file before it can read
    metadata. On a slow connection that routinely exceeded the frontend's
    video-preview timeout even though the download itself had succeeded.
    """
    if shutil.which("ffmpeg") is None:
        return
    temp_path = video_path.with_name(f"{video_path.stem}.faststart{video_path.suffix}")
    try:
        result = subprocess.run(
            [
                "ffmpeg", "-y", "-loglevel", "error",
                "-i", str(video_path),
                "-c", "copy", "-movflags", "+faststart",
                str(temp_path),
            ],
            capture_output=True,
            timeout=300,
        )
        if result.returncode == 0 and temp_path.is_file() and temp_path.stat().st_size > 0:
            temp_path.replace(video_path)
        else:
            print(f"faststart remux skipped for {video_path.name}: {result.stderr.decode('utf-8', 'ignore')[:500]}")
    except Exception as error:
        print(f"faststart remux failed for {video_path.name}: {error}")
    finally:
        if temp_path.exists():
            temp_path.unlink()


def _download_youtube_video_sync(video_id: str, quality: int) -> None:
    """Blocking yt-dlp download and faststart remux, run off the event loop."""
    ydl_opts = {
        'format': (
            f'bestvideo[height<={quality}][vcodec^=vp9]+bestaudio/'
            f'bestvideo[height<={quality}]+bestaudio/best[height<={quality}]'
        ),
        'outtmpl': str(VIDEO_DIR / f'{video_id}.%(ext)s'),
        'merge_output_format': 'mp4',
    }
    cookies_file = os.getenv("YTDLP_COOKIES_FILE")
    if cookies_file:
        cookies_path = Path(cookies_file).expanduser().resolve()
        if not cookies_path.is_file():
            raise RuntimeError(
                f"YTDLP_COOKIES_FILE does not exist: {cookies_path}"
            )
        ydl_opts["cookiefile"] = str(cookies_path)
    with yt_dlp.YoutubeDL(ydl_opts) as ydl:
        try:
            ydl.download([f'https://www.youtube.com/watch?v={video_id}'])
        except yt_dlp.utils.DownloadError as error:
            error_text = str(error)
            if "Sign in to confirm you're not a bot" in error_text:
                raise RuntimeError(
                    "YouTube requires authentication for this download. "
                    "Export YouTube cookies in Netscape format, upload the "
                    "cookie file to Colab, and set YTDLP_COOKIES_FILE to its path."
                ) from error
            raise

    merged_path = VIDEO_DIR / f"{video_id}.mp4"
    if merged_path.is_file():
        _move_moov_atom_to_front(merged_path)
        (VIDEO_DIR / f"{video_id}_video_quality.json").write_text(
            json.dumps({"height": quality}), encoding="utf-8"
        )


@app.get("/download/{video_id}")
@app.post("/download/{video_id}")
async def download_video(video_id: str, background_tasks: BackgroundTasks, quality: int = 480):
    try:
        if quality not in {480, 720}:
            raise HTTPException(status_code=400, detail="Video quality must be 480 or 720")
        video_path = str(VIDEO_DIR / f"{video_id}.mp4")
        quality_path = VIDEO_DIR / f"{video_id}_video_quality.json"
        stored_quality = None
        if quality_path.exists():
            try:
                stored_quality = int(json.loads(quality_path.read_text(encoding="utf-8")).get("height", 0))
            except (OSError, ValueError, TypeError, json.JSONDecodeError):
                stored_quality = None
        existing_video_files = [
            Path(path) for path in glob.glob(str(VIDEO_DIR / f"{video_id}.*"))
            if not path.endswith("_video_quality.json")
        ]
        if existing_video_files and stored_quality and stored_quality < quality:
            for existing_path in existing_video_files:
                existing_path.unlink()
            existing_video_files = []
        if not existing_video_files:
            # Downloading (network I/O) and remuxing (CPU) are blocking calls.
            # Run them in a worker thread so the event loop stays free to serve
            # other requests -- including the browser's request for this
            # video's own preview once this response is sent.
            await asyncio.to_thread(_download_youtube_video_sync, video_id, quality)

        downloaded_files = os.listdir(VIDEO_DIR)
        for file in downloaded_files:
            if file.startswith(video_id) and not file.endswith("_video_quality.json"):
                video_path = str(VIDEO_DIR / file)
                break

        whisper_transcript_path = str(TRANSCRIPTS_DIR / f"{video_id}_whisper.json")
        has_whisper_transcript = os.path.exists(whisper_transcript_path)

        scenes_path = str(SCENES_DIR / f"{video_id}.json")
        has_scenes = False
        if os.path.exists(scenes_path):
            try:
                with open(scenes_path, 'r', encoding='utf-8') as file:
                    has_scenes = bool(json.load(file))
            except (OSError, json.JSONDecodeError):
                has_scenes = False
        
        # Load existing scenes if available
        existing_scenes = []
        if has_scenes:
            try:
                with open(scenes_path, 'r') as f:
                    existing_scenes = json.load(f)
            except Exception as e:
                print(f"Error loading existing scenes: {e}")
        
        # Check if transcript is being generated
        progress_file = str(TRANSCRIPTS_DIR / f"{video_id}_whisper_progress.txt")
        transcript_in_progress = os.path.exists(progress_file)
        
        # Load existing transcript if available
        transcript_to_use = None
        if has_whisper_transcript:
            try:
                with open(whisper_transcript_path, 'r') as f:
                    transcript_to_use = json.load(f)
            except Exception as e:
                print(f"Error loading existing transcript: {e}")
        
        # Try to get YouTube transcript if Whisper transcript is not available
        has_youtube_transcript = False
        if not transcript_to_use and not transcript_in_progress:
            try:
                youtube_transcript = await asyncio.to_thread(
                    YouTubeTranscriptApi.get_transcript, video_id, languages=['en']
                )
                if youtube_transcript:
                    transcript_to_use = youtube_transcript
                    has_youtube_transcript = True
                    with open(str(TRANSCRIPTS_DIR / f"{video_id}_youtube.json"), 'w') as f:
                        json.dump(youtube_transcript, f)
            except Exception as e:
                print(f"Error getting YouTube transcript: {e}")
        
        # Start missing processing tasks
        if not has_whisper_transcript and not transcript_in_progress:
            await transcript_processor.start_whisper_generation(video_id, video_path, background_tasks)
            transcript_in_progress = True
        
        return JSONResponse({
            "success": True,
            "video_url": f"/video/{video_id}",
            "video_id": video_id,
            "transcript": transcript_to_use,
            "has_youtube_transcript": has_youtube_transcript,
            "has_whisper_transcript": has_whisper_transcript,
            "transcript_in_progress": transcript_in_progress,
            "scenes": existing_scenes,
            "is_duplicate": os.path.exists(video_path)
        })
            
    except Exception as e:
        return JSONResponse({
            "success": False,
            "error": str(e)
        })

@app.post("/process_youtube")
async def process_youtube(request: Request, background_tasks: BackgroundTasks):
    try:
        data = await request.json()
        url = data.get("url")
        if not url:
            return JSONResponse({
                "success": False,
                "error": "No URL provided"
            })
        
        video_id = extract_video_id(url)
        if not video_id:
            return JSONResponse({
                "success": False,
                "error": "Invalid YouTube URL"
            })
        
        return await download_video(video_id, background_tasks)
        
    except Exception as e:
        return JSONResponse({
            "success": False,
            "error": str(e)
        })

@app.post("/generate_summary")
async def generate_summary(request: Request):
    try:
        data = await request.json()
        transcript = data.get("transcript", [])
        video_id = data.get("video_id")
        return await summary_processor.generate_summary(transcript, video_id, data.get("model"))
    except Exception as e:
        return JSONResponse({
            "success": False,
            "error": str(e)
        })

@app.get("/summary/{video_id}")
async def get_summary(video_id: str):
    return await summary_processor.get_summary(video_id)

@app.get("/export_suggestions/{video_id}")
async def get_export_suggestions(video_id: str):
    metadata_path = VIDEO_DIR / "metadata.json"
    source_title = video_id
    if metadata_path.is_file():
        try:
            metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
            source_title = Path(
                metadata.get(video_id, {}).get("filename", source_title)
            ).stem or source_title
        except (OSError, json.JSONDecodeError):
            pass

    title_suggestion = source_title
    summary_path = SUMMARIES_DIR / f"{video_id}.json"
    if summary_path.is_file():
        try:
            chapters = json.loads(summary_path.read_text(encoding="utf-8"))
            title_suggestion = next(
                (
                    str(chapter.get("title", "")).strip()
                    for chapter in chapters
                    if isinstance(chapter, dict)
                    and str(chapter.get("title", "")).strip()
                    and not str(chapter.get("title", "")).lower().startswith(("part ", "chapter "))
                ),
                source_title,
            )
        except (OSError, json.JSONDecodeError, TypeError):
            pass

    filename_suggestion = re.sub(
        r"[^A-Za-z0-9._-]+", "_", title_suggestion
    ).strip("._-") or f"video_{video_id}"
    return {
        "title_suggestion": title_suggestion,
        "filename_suggestion": filename_suggestion,
        "source_filename": source_title,
    }

def parse_chapter_timestamp(timestamp: str) -> float:
    parts = [float(part) for part in timestamp.split(":")]
    if len(parts) == 2:
        return parts[0] * 60 + parts[1]
    if len(parts) == 3:
        return parts[0] * 3600 + parts[1] * 60 + parts[2]
    raise ValueError(f"Invalid chapter timestamp: {timestamp}")

def format_chapter_timestamp(seconds: float) -> str:
    total_seconds = max(0, int(seconds))
    hours, remainder = divmod(total_seconds, 3600)
    minutes, seconds = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def chapter_boundaries_from_topics(chapters):
    return [
        {"timestamp": str(chapter["timestamp"]), "title": str(chapter.get("title", f"Chapter {index + 1}"))}
        for index, chapter in enumerate(chapters)
        if isinstance(chapter, dict) and "timestamp" in chapter
    ]


def chapter_boundaries_from_scenes(scenes):
    return [
        {"timestamp": str(scene["timestamp"]), "title": f"Slide change {index + 1}"}
        for index, scene in enumerate(scenes)
        if isinstance(scene, dict) and "timestamp" in scene
    ]


def combine_chapter_boundaries(topic_chapters, scene_chapters):
    boundaries = {}
    for chapter in topic_chapters + scene_chapters:
        seconds = parse_chapter_timestamp(chapter["timestamp"])
        existing = boundaries.get(seconds)
        if existing is None or existing.startswith("Slide change"):
            boundaries[seconds] = chapter["title"]
    return [
        {"timestamp": format_chapter_timestamp(seconds), "title": title}
        for seconds, title in sorted(boundaries.items())
    ]

def create_combined_chapter_documents(
    chapters, chapter_files, pdf_path=None, docx_path=None,
    pdf_text=True, pdf_images=True, word_text=True, word_images=True,
    document_title="Lecture Chapters", document_subtitle=None
):
    """Create title, screenshot, and transcript documents for all chapters."""
    styles = getSampleStyleSheet()
    pdf_story = [Paragraph(document_title, styles["Title"])] if pdf_path else None
    if pdf_story is not None and document_subtitle:
        pdf_story.append(Paragraph(document_subtitle, styles["Italic"]))
    word_document = Document() if docx_path else None
    if word_document:
        word_document.add_heading(document_title, level=0)
        if document_subtitle:
            word_document.add_paragraph(document_subtitle, style="Subtitle")

    for index, chapter in enumerate(chapters):
        files = chapter_files[index]
        title = str(chapter["title"])
        timestamp = format_chapter_timestamp(chapter["start"])
        transcript = files["transcript"].read_text(encoding="utf-8").strip()
        heading = f"Part {index + 1}: {title}"
        image_path = files["image"]
        try:
            with PillowImage.open(image_path) as image:
                image.verify()
            normalized_image_path = image_path.with_suffix(".png")
            with PillowImage.open(image_path) as image:
                image.convert("RGB").save(normalized_image_path, format="PNG")
            image_path = normalized_image_path
        except (FileNotFoundError, UnidentifiedImageError, OSError):
            image_path = None

        if pdf_story is not None:
            pdf_story.extend([
            Paragraph(heading, styles["Heading1"]),
            Paragraph(f"Starts at {timestamp}", styles["Normal"]),
            Spacer(1, 0.15 * inch),
            ])
        if pdf_story is not None and pdf_images and image_path is not None:
            pdf_story.extend([
                PdfImage(str(image_path), width=6.5 * inch, height=3.65 * inch, kind="proportional"),
                Spacer(1, 0.15 * inch),
            ])
        if pdf_story is not None and pdf_text:
            pdf_story.append(Paragraph(
            (transcript or "No transcript available for this chapter.")
            .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace("\n", "<br/>"),
            styles["BodyText"],
            ))
        if word_document:
            word_document.add_heading(heading, level=1)
            word_document.add_paragraph(f"Starts at {timestamp}")
        if word_document and word_images and image_path is not None:
            try:
                word_document.add_picture(str(image_path), width=Inches(6.5))
            except UnrecognizedImageError:
                pass
        if word_document and word_text:
            word_document.add_heading("Transcript", level=2)
            word_document.add_paragraph(transcript or "No transcript available for this chapter.")
    footnote = "Transcription model: faster-whisper turbo."
    if pdf_story is not None:
        pdf_story.extend([Spacer(1, 0.3 * inch), Paragraph(footnote, styles["Italic"])])
        SimpleDocTemplate(str(pdf_path), pagesize=letter, rightMargin=0.6 * inch,
                      leftMargin=0.6 * inch, topMargin=0.6 * inch,
                      bottomMargin=0.6 * inch).build(pdf_story)
    if word_document:
        word_document.add_paragraph(footnote, style="Caption")
        word_document.save(str(docx_path))

def collapse_word_heading(paragraph):
    paragraph_properties = paragraph._p.get_or_add_pPr()
    paragraph_properties.append(OxmlElement("w:collapsed"))

@app.post("/export_chapters/{video_id}")
async def export_chapters(video_id: str, request: Request):
    """Create a ZIP containing chapter clips, screenshots, and transcripts."""
    options = await request.json()
    interval_minutes = options.get("interval_minutes")
    transcript_language = str(options.get("transcript_language", "")).strip()
    if transcript_language and transcript_language not in {"de", "en", "ar", "pl"}:
        raise HTTPException(status_code=400, detail="Invalid transcript language")
    chapter_grouping = options.get("chapter_grouping", "topic")
    if chapter_grouping not in {"topic", "slides", "combined"}:
        raise HTTPException(status_code=400, detail="Invalid chapter grouping")
    timestamp_mode = options.get("timestamp_mode", "original")
    if timestamp_mode not in {"original", "part", "subpart"}:
        raise HTTPException(status_code=400, detail="Invalid timestamp mode")
    subpart_mode = options.get("subpart_mode", "points")
    if subpart_mode not in {"points", "slides", "both"}:
        raise HTTPException(status_code=400, detail="Invalid subpart mode")
    export_flags = {name: bool(options.get(name, True)) for name in (
        "include_images", "include_transcripts", "include_clips",
        "include_word", "include_pdf", "include_webpage", "include_outline", "include_scorm",
    )}
    if not any(export_flags.values()):
        raise HTTPException(status_code=400, detail="Select at least one export option")
    video_path = video_processor.get_video_path(video_id)
    translated_summary_path = (
        SUMMARIES_DIR / f"{video_id}_summary_{transcript_language}.json"
        if transcript_language else None
    )
    summary_path = translated_summary_path if translated_summary_path and translated_summary_path.is_file() else SUMMARIES_DIR / f"{video_id}.json"
    if not video_path:
        raise HTTPException(status_code=404, detail="Video not found")
    scene_path = SCENES_DIR / f"{video_id}.json"
    if interval_minutes is None and chapter_grouping in {"slides", "combined"} and not scene_path.is_file():
        if chapter_grouping == "slides":
            raise HTTPException(status_code=400, detail="Detect slides before exporting by slide changes")

    try:
        scene_chapters = []
        if interval_minutes is not None:
            try:
                interval_seconds = float(interval_minutes) * 60
            except (TypeError, ValueError):
                raise HTTPException(status_code=400, detail="Interval duration must be a number of minutes")
            if interval_seconds <= 0:
                raise HTTPException(status_code=400, detail="Interval duration must be greater than zero")
            probe = subprocess.run(
                ["ffprobe", "-v", "error", "-show_entries", "format=duration",
                 "-of", "default=noprint_wrappers=1:nokey=1", str(video_path)],
                check=True, capture_output=True, text=True
            )
            duration = float(probe.stdout.strip())
            chapters = [
                {"timestamp": format_chapter_timestamp(start), "title": f"Interval {index + 1}"}
                for index, start in enumerate(
                    range(0, max(1, int(duration)), max(1, int(interval_seconds)))
                )
            ]
        else:
            topic_chapters = []
            scene_chapters = []
            if summary_path.is_file():
                with summary_path.open("r", encoding="utf-8") as file:
                    topic_chapters = chapter_boundaries_from_topics(json.load(file))
            if scene_path.is_file():
                with scene_path.open("r", encoding="utf-8") as file:
                    scene_chapters = chapter_boundaries_from_scenes(json.load(file))
            if chapter_grouping == "topic":
                chapters = topic_chapters
            elif chapter_grouping == "slides":
                chapters = scene_chapters
            else:
                # Keep topics as the top-level parts so slide changes can be
                # represented as subparts in the outline.
                chapters = topic_chapters
        if not isinstance(chapters, list) or not chapters:
            if chapter_grouping == "slides":
                raise HTTPException(status_code=400, detail="Detect slides before exporting by slide changes")
            # Keep asset and transcript exports usable when AI chapter
            # generation failed. The full video becomes one export part.
            chapters = [{"timestamp": "00:00", "title": "Full video"}]

        transcript = []
        transcript_candidates = (
            [TRANSCRIPTS_DIR / f"{video_id}_translated_{transcript_language}.json"]
            if transcript_language else []
        ) + [
            TRANSCRIPTS_DIR / f"{video_id}_whisper.json",
            TRANSCRIPTS_DIR / f"{video_id}_youtube.json",
            TRANSCRIPTS_DIR / f"{video_id}.json",
        ]
        for transcript_path in transcript_candidates:
            if transcript_path.is_file():
                with transcript_path.open("r", encoding="utf-8") as file:
                    transcript = json.load(file)
                break
        transcript_required = any((
            export_flags["include_transcripts"],
            export_flags["include_word"],
            export_flags["include_pdf"],
            export_flags["include_webpage"],
            export_flags["include_outline"],
            export_flags["include_scorm"],
        ))
        if not transcript and transcript_required:
            raise HTTPException(status_code=400, detail="No transcript available")
        speaker_names = {}
        speaker_meta_path = TRANSCRIPTS_DIR / f"{video_id}_whisper_meta.json"
        if speaker_meta_path.is_file():
            try:
                speaker_names = json.loads(
                    speaker_meta_path.read_text(encoding="utf-8")
                ).get("speaker_names", {})
            except (OSError, json.JSONDecodeError):
                speaker_names = {}

        chapter_data = []
        for index, chapter in enumerate(chapters):
            if not isinstance(chapter, dict) or "timestamp" not in chapter:
                continue
            start = parse_chapter_timestamp(str(chapter["timestamp"]))
            end = (
                parse_chapter_timestamp(str(chapters[index + 1]["timestamp"]))
                if index + 1 < len(chapters)
                else None
            )
            chapter_data.append({
                "index": len(chapter_data) + 1,
                "title": str(chapter.get("title", f"Chapter {index + 1}")),
                "start": start,
                "end": end,
            })
        if not chapter_data:
            raise HTTPException(status_code=400, detail="Chapter timestamps are invalid")

        export_dir = EXPORTS_DIR / video_id
        export_dir.mkdir(parents=True, exist_ok=True)
        zip_path = EXPORTS_DIR / f"{video_id}_chapters.zip"
        direct_export_path = None
        document_title = Path(video_path).stem
        source_filename = document_title
        metadata_path = VIDEO_DIR / "metadata.json"
        if metadata_path.is_file():
            try:
                metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
                document_title = Path(
                    metadata.get(video_id, {}).get("filename", document_title)
                ).stem or document_title
            except (OSError, json.JSONDecodeError):
                pass
        requested_title = str(options.get("document_title", "")).strip()
        if requested_title:
            document_title = requested_title
        with tempfile.TemporaryDirectory(dir=EXPORTS_DIR) as temp_dir_name:
            temp_dir = Path(temp_dir_name)
            generic_name = (
                bool(re.fullmatch(r"[0-9a-f]{32,64}", document_title, re.IGNORECASE))
                or document_title.lower() in {"video", "meeting", "recording", "input"}
            )
            topic_title = next(
                (
                    str(chapter.get("title", "")).strip()
                    for chapter in chapter_data
                    if str(chapter.get("title", "")).strip()
                    and not str(chapter.get("title", "")).lower().startswith(("part ", "chapter "))
                ),
                "",
            )
            document_subtitle = f"Source clip: {source_filename}" if topic_title and generic_name else None
            if topic_title and generic_name:
                document_title = topic_title
            export_title = re.sub(r"[^A-Za-z0-9._-]+", "_", document_title).strip("._-")
            export_title = export_title or f"video_{video_id}"
            requested_filename = re.sub(
                r"[^A-Za-z0-9._-]+", "_", str(options.get("export_filename", "")).strip()
            ).strip("._-")
            zip_path = EXPORTS_DIR / (
                requested_filename if requested_filename.lower().endswith(".zip")
                else requested_filename + ".zip"
            ) if requested_filename else EXPORTS_DIR / f"{export_title}_chapters.zip"
            chaptered_lines = [f"# {document_title}", ""]
            chapter_files = []
            for chapter in chapter_data:
                number = chapter["index"]
                safe_title = re.sub(r"[^A-Za-z0-9_-]+", "_", chapter["title"]).strip("_") or f"chapter_{number}"
                base_name = f"{number:02d}_{safe_title}"
                clip_path = temp_dir / f"{base_name}.mp4"
                image_path = temp_dir / f"{base_name}.jpg"
                transcript_path = temp_dir / f"{base_name}.txt"

                ffmpeg_clip = [
                    "ffmpeg", "-y", "-hide_banner", "-loglevel", "error",
                    "-ss", str(chapter["start"]),
                    "-i", str(video_path),
                ]
                if chapter["end"] is not None:
                    ffmpeg_clip.extend(["-t", str(chapter["end"] - chapter["start"])])
                ffmpeg_clip.extend(["-c:v", "libx264", "-c:a", "aac", str(clip_path)])
                if export_flags["include_clips"] or export_flags["include_webpage"] or export_flags["include_scorm"]:
                    subprocess.run(ffmpeg_clip, check=True)
                if export_flags["include_images"] or export_flags["include_word"] or export_flags["include_pdf"] or export_flags["include_webpage"] or export_flags["include_scorm"]:
                    subprocess.run([
                    "ffmpeg", "-y", "-hide_banner", "-loglevel", "error",
                    "-ss", str(chapter["start"]), "-i", str(video_path),
                    "-frames:v", "1", str(image_path)
                    ], check=True)

                chapter_transcript = [
                    item for item in transcript
                    if chapter["start"] <= float(item.get("start", 0))
                    and (chapter["end"] is None or float(item.get("start", 0)) < chapter["end"])
                ]
                def format_export_transcript_item(item):
                    speaker = item.get("speaker")
                    speaker_prefix = (
                        f"[{speaker_names.get(speaker, speaker)}] " if speaker else ""
                    )
                    return (
                        f"[{format_chapter_timestamp(float(item.get('start', 0)))}] "
                        f"{speaker_prefix}{item.get('text', '').strip()}"
                    )

                transcript_text = "\n".join(
                    format_export_transcript_item(item) for item in chapter_transcript
                )
                if timestamp_mode != "original":
                    transcript_lines = transcript_text.splitlines()
                    if timestamp_mode == "part":
                        transcript_lines = [
                            re.sub(r"^\[[^\]]+\]\s*", "", line)
                            for line in transcript_lines
                        ]
                    else:
                        transcript_lines = [
                            line if index == 0 else re.sub(r"^\[[^\]]+\]\s*", "", line)
                            for index, line in enumerate(transcript_lines)
                        ]
                    transcript_text = "\n".join(transcript_lines)
                if export_flags["include_transcripts"] or export_flags["include_word"] or export_flags["include_pdf"] or export_flags["include_webpage"] or export_flags["include_scorm"] or export_flags["include_outline"]:
                    transcript_path.write_text(transcript_text + "\n", encoding="utf-8")
                chapter_files.append({
                    "image": image_path,
                    "clip": clip_path,
                    "transcript": transcript_path,
                    "outline_points": build_outline_points(chapter_transcript),
                    "slide_subparts": [
                        scene for scene in scene_chapters
                        if chapter["start"] <= parse_chapter_timestamp(scene["timestamp"])
                        and (chapter["end"] is None or parse_chapter_timestamp(scene["timestamp"]) < chapter["end"])
                    ],
                })
                chaptered_lines.extend([
                    f"## Part {number}: {chapter['title']}",
                    f"Start: {format_chapter_timestamp(chapter['start'])}",
                    "",
                    transcript_text,
                    "",
                ])

            if export_flags["include_transcripts"]:
                (temp_dir / "transcript_by_chapter.md").write_text("\n".join(chaptered_lines), encoding="utf-8")
            (temp_dir / "chapters.json").write_text(json.dumps(chapter_data, indent=2), encoding="utf-8")
            create_combined_chapter_documents(
                chapter_data, chapter_files,
                temp_dir / "chapter_document.pdf" if export_flags["include_pdf"] else None,
                temp_dir / "chapter_document.docx" if export_flags["include_word"] else None,
                pdf_text=export_flags["include_pdf"], pdf_images=export_flags["include_pdf"],
                word_text=export_flags["include_word"], word_images=export_flags["include_word"],
                document_title=document_title,
                document_subtitle=document_subtitle,
            )
            if export_flags["include_webpage"] or export_flags["include_scorm"]:
                webpage_parts = [
                    "<!doctype html><html lang=\"en\"><head><meta charset=\"utf-8\">",
                    "<meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">",
                    f"<title>{html.escape(document_title)}</title>",
                    "<style>body{font-family:Arial,sans-serif;max-width:1000px;margin:2rem auto;padding:0 1rem;color:#222}"
                    "section{border-top:1px solid #ccc;padding:2rem 0}img,video{max-width:100%;display:block;margin:1rem 0}"
                    "pre{white-space:pre-wrap;background:#f6f6f6;padding:1rem;border-radius:6px}.subtitle{color:#666}</style></head><body>",
                    f"<h1>{html.escape(document_title)}</h1>",
                ]
                if export_flags["include_scorm"]:
                    webpage_parts.insert(1, "<script src=\"scorm_api.js\"></script>")
                if document_subtitle:
                    webpage_parts.append(f"<p class=\"subtitle\">{html.escape(document_subtitle)}</p>")
                for chapter, files in zip(chapter_data, chapter_files):
                    image_name = files["image"].name
                    clip_name = files["clip"].name
                    transcript_html = html.escape(files["transcript"].read_text(encoding="utf-8"))
                    point_details = "".join(
                        f"<details><summary>Key point — {format_chapter_timestamp(point['start'])}</summary>"
                        f"<p>{html.escape(point['text'])}</p></details>"
                        for point in files["outline_points"]
                    )
                    if not point_details:
                        point_details = "<p>No key points were available for this part.</p>"
                    webpage_parts.extend([
                        f"<section><h2>Part {chapter['index']}: {html.escape(chapter['title'])}</h2>",
                        f"<p>Start: {format_chapter_timestamp(chapter['start'])}</p>",
                            f"<details><summary><strong>Part {chapter['index']}: {html.escape(chapter['title'])}</strong> "
                            f"({format_chapter_timestamp(chapter['start'])})</summary>",
                            f"<img src=\"{html.escape(image_name)}\" alt=\"Slide for part {chapter['index']}\">",
                            f"<video controls preload=\"metadata\" src=\"{html.escape(clip_name)}\"></video>",
                            f"<details><summary>Transcript and key points</summary>{point_details}"
                            f"<pre>{transcript_html}</pre></details>",
                            "</details></section>",
                        ])
                webpage_parts.append("</body></html>")
                (temp_dir / "index.html").write_text("\n".join(webpage_parts), encoding="utf-8")
            if export_flags["include_outline"]:
                outline_document = Document()
                outline_document.add_heading(document_title, level=0)
                if document_subtitle:
                        outline_document.add_paragraph(document_subtitle, style="Subtitle")
                outline_document.add_heading("Video outline", level=1)
                for chapter, files in zip(chapter_data, chapter_files):
                        outline_document.add_heading(
                            f"Part {chapter['index']}: {chapter['title']}", level=2
                        )
                        outline_document.add_paragraph(
                            f"Timestamp: {format_chapter_timestamp(chapter['start'])}"
                        )
                        outline_document.add_paragraph(
                            "Key points",
                            style="Intense Quote",
                        )
                        outline_items = []
                        if subpart_mode in {"points", "both"}:
                            outline_items.extend(
                                ("point", point["start"], point["title"], point["text"])
                                for point in files["outline_points"]
                            )
                        if subpart_mode in {"slides", "both"}:
                            outline_items.extend(
                                (
                                    "slide",
                                    parse_chapter_timestamp(scene["timestamp"]),
                                    scene["title"],
                                    scene["title"],
                                )
                                for scene in files["slide_subparts"]
                            )
                        outline_items.sort(key=lambda item: item[1])
                        if outline_items:
                            for label, start, title, text in outline_items:
                                outline_document.add_heading(
                                    f"{title} ({format_chapter_timestamp(start)})",
                                    level=3,
                                )
                                if label == "point":
                                    outline_document.add_paragraph(
                                        text,
                                        style="List Bullet 2",
                                    )
                        else:
                            outline_document.add_paragraph(
                                "No spoken content was available for this part.",
                                style="List Bullet 2",
                            )
                        supporting_heading = outline_document.add_heading(
                            "Picture and transcript",
                            level=3,
                        )
                        collapse_word_heading(supporting_heading)
                        image_path = files["image"]
                        if image_path.is_file():
                            try:
                                outline_document.add_picture(
                                    str(image_path), width=Inches(6.5)
                                )
                            except UnrecognizedImageError:
                                pass
                        outline_document.add_heading("Transcript", level=4)
                        outline_document.add_paragraph(
                            files["transcript"].read_text(encoding="utf-8").strip()
                            or "No transcript available for this part."
                        )
                outline_document.save(str(temp_dir / "video_outline.docx"))
            if export_flags["include_scorm"]:
                (temp_dir / "scorm_api.js").write_text(
                        "window.addEventListener('load',()=>{"
                        "try{if(window.parent&&window.parent.API){window.parent.API.LMSInitialize('');"
                        "window.parent.API.LMSSetValue('cmi.core.lesson_status','completed');"
                        "window.parent.API.LMSCommit('');}}catch(e){console.warn('SCORM API unavailable',e);}});",
                        encoding="utf-8",
                )
                (temp_dir / "imsmanifest.xml").write_text(
                        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
                        "<manifest identifier=\"slidedec-video\" version=\"1.2\" "
                        "xmlns=\"http://www.imsproject.org/xsd/imscp_rootv1p1p2\" "
                        "xmlns:adlcp=\"http://www.adlnet.org/xsd/adlcp_rootv1p2\" "
                        "xmlns:xsi=\"http://www.w3.org/2001/XMLSchema-instance\" "
                        "xsi:schemaLocation=\"http://www.imsproject.org/xsd/imscp_rootv1p1p2 "
                        "http://www.imsglobal.org/xsd/imscp_rootv1p1p2.xsd "
                        "http://www.adlnet.org/xsd/adlcp_rootv1p2 "
                        "http://www.imsglobal.org/xsd/adlcp_rootv1p2.xsd\">"
                        "<organizations default=\"org1\"><organization identifier=\"org1\">"
                        f"<title>{html.escape(document_title)}</title><item identifier=\"item1\" "
                        "identifierref=\"resource1\"><title>Video lecture</title></item>"
                        "</organization></organizations><resources><resource identifier=\"resource1\" "
                        "type=\"webcontent\" adlcp:scormtype=\"sco\" href=\"index.html\">"
                        "<file href=\"index.html\"/><file href=\"scorm_api.js\"/>"
                        + "".join(
                            f"<file href=\"{html.escape(path.name)}\"/>"
                            for path in temp_dir.iterdir()
                            if path.name not in {"imsmanifest.xml", "index.html", "scorm_api.js"}
                        )
                        + "</resource></resources></manifest>",
                        encoding="utf-8",
                )
            only_word = export_flags["include_word"] and not export_flags["include_pdf"]
            only_pdf = export_flags["include_pdf"] and not export_flags["include_word"]
            only_outline = export_flags["include_outline"] and not any(
                export_flags[name] for name in (
                    "include_word", "include_pdf", "include_webpage", "include_scorm",
                    "include_images", "include_transcripts", "include_clips",
                )
            )
            no_extra_files = not any(
                export_flags[name] for name in (
                    "include_images", "include_transcripts", "include_clips",
                    "include_webpage", "include_outline", "include_scorm",
                )
            )
            if no_extra_files and (only_word or only_pdf):
                document_name = "chapter_document.docx" if only_word else "chapter_document.pdf"
                requested_filename = str(options.get("export_filename", "")).strip()
                requested_filename = re.sub(r"[^A-Za-z0-9._-]+", "_", requested_filename).strip("._-")
                if requested_filename:
                    suffix = ".docx" if only_word else ".pdf"
                    direct_export_path = EXPORTS_DIR / (
                        requested_filename if requested_filename.lower().endswith(suffix)
                        else requested_filename + suffix
                    )
                else:
                    direct_export_path = EXPORTS_DIR / f"{export_title}_{'word' if only_word else 'pdf'}.{'docx' if only_word else 'pdf'}"
                shutil.copy2(temp_dir / document_name, direct_export_path)
            elif only_outline:
                requested_filename = str(options.get("export_filename", "")).strip()
                requested_filename = re.sub(r"[^A-Za-z0-9._-]+", "_", requested_filename).strip("._-")
                direct_export_path = EXPORTS_DIR / (
                    (requested_filename if requested_filename.lower().endswith(".docx") else requested_filename + ".docx")
                    if requested_filename else f"{export_title}_outline.docx"
                )
                shutil.copy2(temp_dir / "video_outline.docx", direct_export_path)
            with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
                for file_path in temp_dir.iterdir():
                    if file_path.suffix == ".mp4" and not (export_flags["include_clips"] or export_flags["include_webpage"] or export_flags["include_scorm"]):
                        continue
                    if file_path.suffix.lower() in {".jpg", ".jpeg", ".png"} and not (export_flags["include_images"] or export_flags["include_webpage"] or export_flags["include_scorm"]):
                        continue
                    if file_path.suffix == ".txt" and not (export_flags["include_transcripts"] or export_flags["include_webpage"] or export_flags["include_scorm"]):
                        continue
                    archive.write(file_path, file_path.name)

        if direct_export_path is not None:
            return FileResponse(
                direct_export_path,
                media_type=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    if direct_export_path.suffix == ".docx" else "application/pdf"
                ),
                filename=direct_export_path.name,
            )
        return FileResponse(
            zip_path,
            media_type="application/zip",
            filename=zip_path.name
        )
    except HTTPException:
        raise
    except (ValueError, OSError, subprocess.CalledProcessError, json.JSONDecodeError) as error:
        raise HTTPException(status_code=500, detail=f"Chapter export failed: {error}")
    except Exception as error:
        error_type = type(error).__name__
        traceback.print_exc()
        raise HTTPException(
            status_code=500,
            detail=f"Chapter export failed unexpectedly ({error_type}): {error!r}"
        )

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)